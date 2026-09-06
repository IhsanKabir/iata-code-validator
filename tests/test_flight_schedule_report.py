"""The Flight Schedule History workbook."""
from datetime import datetime

from openpyxl import load_workbook

from src import flight_schedule_history as fsh
from src import flight_schedule_report as fsr
from tests.test_flight_schedule_history import (TRANSFER, _Event, _Flight,
                                                _time_change)


def _text(ws):
    return " | ".join(str(c.value) for row in ws.iter_rows() for c in row
                      if c.value)


def _built(tmp_path, events, roster=None, month=8, year=2026):
    res = fsh.aggregate(events, roster=roster)
    out = fsr.build_workbook(res, tmp_path / "fsh.xlsx", month=month,
                             year=year, roster=roster)
    return res, load_workbook(out, data_only=True)


class _Ref:
    def __init__(self, number, date_, o, d, dep="", arr="", ac=""):
        self.flight_number, self.flight_date = number, date_
        self.origin, self.destination = o, d
        self.sched_dep, self.sched_arr, self.aircraft = dep, arr, ac


FLIGHT = _Flight("BS381", "DAC", "RUH", "28/07/2026")


def test_the_workbook_has_the_sheets_the_question_needs(tmp_path):
    roster = [_Ref("BS381", "28/07/2026", "DAC", "RUH", "12:55", "17:10", "738")]
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0), pnr="0A1",
               flight=FLIGHT),
        _Event("seat", when=datetime(2026, 7, 1, 9, 0), etype="Seat map change"),
    ], roster=roster)
    assert wb.sheetnames[:4] == ["Summary", "By route", "By flight",
                                 "Every change"]
    assert "Schedule" in wb.sheetnames
    assert "Unrecognised" in wb.sheetnames


def test_a_swaps_date_gap_is_never_printed_as_a_delay(tmp_path):
    """47,520 minutes in a column headed 'Shift (min)' would be read as a
    delay. A swap moved the flight to another date; it did not retime it."""
    _res, wb = _built(tmp_path, [
        _Event(TRANSFER, when=datetime(2026, 7, 2, 9, 0), pnr="0A1",
               flight=FLIGHT, etype="Flight transfer")])
    ws = wb["Every change"]
    header = [c.value for c in ws[4]]
    shift_col = header.index("Shift (min)") + 1
    body = [ws.cell(row=r, column=shift_col).value
            for r in range(5, ws.max_row + 1)]
    assert all(v is None for v in body)
    assert "Swapped / replaced" in _text(ws)


def test_a_time_change_does_print_its_shift(tmp_path):
    _res, wb = _built(tmp_path, [
        _Event(_time_change(old="12:55", new="13:55"),
               when=datetime(2026, 7, 1, 9, 0), flight=FLIGHT)])
    ws = wb["Every change"]
    header = [c.value for c in ws[4]]
    col = header.index("Shift (min)") + 1
    assert ws.cell(row=5, column=col).value == 60


def test_the_summary_calls_the_passenger_count_a_floor_when_it_is_one(tmp_path):
    """No repetition observed means each change was seen once, so the count
    is a lower bound and must not read as a census."""
    res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0), pnr="0A1",
               flight=FLIGHT)])
    assert not res.repetition_observed
    text = _text(wb["Summary"])
    assert "floor" in text
    assert "Passengers affected is a floor" in text


def test_the_summary_says_so_when_repetition_was_actually_seen(tmp_path):
    when = datetime(2026, 7, 1, 9, 0)
    res, wb = _built(tmp_path, [
        _Event(_time_change(), when=when, pnr=p, flight=FLIGHT)
        for p in ("0A1", "0A2")])
    assert res.repetition_observed
    text = _text(wb["Summary"])
    assert "counted from repeated rows" in text
    assert "Passengers affected is a floor" not in text


def test_the_summary_states_the_missing_roster_rather_than_hiding_it(tmp_path):
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)])
    text = _text(wb["Summary"])
    assert "Share of flights changed needs the roster" in text
    assert "would always read as 100%" in text


def test_share_changed_is_blank_on_the_sheet_without_a_roster(tmp_path):
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)])
    ws = wb["By route"]
    header = [c.value for c in ws[4]]
    col = header.index("Share changed") + 1
    assert ws.cell(row=5, column=col).value is None


def test_the_summary_admits_changes_before_any_booking_are_invisible(tmp_path):
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)])
    text = _text(wb["Summary"])
    assert "before any booking" in text
    assert "lower bound" in text


def test_the_schedule_sheet_carries_the_times_and_aircraft(tmp_path):
    roster = [_Ref("BS381", "28/07/2026", "DAC", "RUH", "12:55", "17:10",
                   "Boeing 737-800")]
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)], roster=roster)
    text = _text(wb["Schedule"])
    assert "12:55" in text and "17:10" in text
    assert "Boeing 737-800" in text


def test_a_blank_scheduled_time_is_not_filled_in(tmp_path):
    roster = [_Ref("BS381", "28/07/2026", "DAC", "RUH")]
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)], roster=roster)
    ws = wb["Schedule"]
    header = [c.value for c in ws[4]]
    col = header.index("Scheduled departure") + 1
    assert ws.cell(row=5, column=col).value is None
    assert "did not carry one" in _text(ws)


def test_an_unrecognised_sheet_appears_only_when_there_is_a_blind_spot(tmp_path):
    _res, wb = _built(tmp_path, [
        _Event(_time_change(), when=datetime(2026, 7, 1, 9, 0),
               flight=FLIGHT)])
    assert "Unrecognised" not in wb.sheetnames


def test_the_route_rollup_orders_the_worst_first(tmp_path):
    quiet = _Flight("BS500", "DAC", "CGP", "28/07/2026")
    events = [
        _Event(_time_change(route="DACCGP"), when=datetime(2026, 7, 1, 9, 0),
               flight=quiet),
    ] + [
        _Event(_time_change(new=f"1{h}:55"),
               when=datetime(2026, 7, d, 9, 0), flight=FLIGHT)
        for d, h in ((1, 3), (2, 4), (3, 5))
    ]
    _res, wb = _built(tmp_path, events)
    ws = wb["By route"]
    assert ws.cell(row=5, column=1).value == "DAC-RUH"      # 3 changes
    assert ws.cell(row=6, column=1).value == "DAC-CGP"      # 1 change


def test_an_empty_result_still_produces_a_readable_workbook(tmp_path):
    _res, wb = _built(tmp_path, [])
    assert "Summary" in wb.sheetnames
    assert "FLIGHT SCHEDULE HISTORY" in _text(wb["Summary"])
