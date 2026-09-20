"""The Sales Movement workbook.

Fixtures are synthetic: the warehouse carries real agency names and revenue.
"""
from datetime import date

from openpyxl import load_workbook

from src import sales_movement as sm
from src import sales_movement_report as smr

AUG = sm.Settings(period_from=date(2026, 8, 1), period_to=date(2026, 8, 31),
                  trailing=3, threshold=0.20, floor=500_000)


def _row(customer, window, amount, **meta):
    base = {"customer": customer, "window": window, "amount": amount,
            "tickets": 10, "customer_id": "", "iata": "1234567",
            "zone": "Zone-4", "station": "DAC", "sales_person": "Rep A",
            "agent_type": "BSP", "channel": "AGENCY"}
    base.update(meta)
    return base


def _result(**kw):
    rows = []
    rows += [_row("Falling", 0, 3_000_000)] + \
        [_row("Falling", i, 10_000_000) for i in (1, 2, 3)]
    rows += [_row("SmallFall", 0, 100_000)] + \
        [_row("SmallFall", i, 900_000) for i in (1, 2, 3)]
    rows += [_row("Rising", 0, 9_000_000)] + \
        [_row("Rising", i, 3_000_000, zone="Zone-9") for i in (1, 2, 3)]
    rows += [_row("Gone", i, 4_000_000, zone="Zone-9") for i in (1, 2, 3)]
    rows += [_row("Fresh", 0, 2_000_000, sales_person="Rep B")]
    return sm.build(rows, AUG, **kw)


def _book(tmp_path, res=None):
    out = tmp_path / "movement.xlsx"
    smr.build_workbook(res or _result(), out)
    return load_workbook(out)


def _column(ws, header, limit=60):
    """Values under a header, so a test never hard-codes a column letter.

    The Summary's tables sit below a KPI strip and any warnings, so their
    row is not fixed -- the header is searched for across the whole sheet.
    """
    for r in range(1, ws.max_row + 1):
        for c in range(1, 22):
            if ws.cell(row=r, column=c).value == header:
                return [ws.cell(row=i, column=c).value
                        for i in range(r + 1, min(r + 1 + limit,
                                                  ws.max_row + 1))]
    raise AssertionError(f"no column headed {header!r}")


def _text(ws) -> str:
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row
                    if c.value is not None)


# --------------------------------------------------------------------------
def test_every_group_gets_its_own_sheet(tmp_path):
    wb = _book(tmp_path)
    assert wb.sheetnames == ["Summary", "Declined", "Stopped buying",
                             "Refunded more than sold", "Grew", "New"]


def test_the_summary_states_the_settings_in_words(tmp_path):
    said = _text(_book(tmp_path)["Summary"])
    assert "20%" in said and "500,000" in said
    assert "Aug 2026" in said and "May 2026" in said


def test_the_summary_lists_the_windows_actually_compared(tmp_path):
    ws = _book(tmp_path)["Summary"]
    roles = [v for v in _column(ws, "Window") if v]
    assert roles[:2] == ["This period", "Baseline 1"]


def test_decliners_are_ranked_by_money_not_by_percentage(tmp_path):
    """SmallFall is down 89%, Falling only 70% -- but Falling lost 7,000,000
    and SmallFall lost 800,000, so Falling belongs on top."""
    names = [v for v in _column(_book(tmp_path)["Declined"],
                                "Agency / customer") if v]
    assert names == ["Falling", "SmallFall"]


def test_a_new_customer_gets_no_percentage_cell(tmp_path):
    ws = _book(tmp_path)["New"]
    assert [v for v in _column(ws, "Agency / customer") if v] == ["Fresh"]
    assert _column(ws, "Change %")[0] is None


def test_a_lapsed_customer_is_on_its_own_sheet_with_the_baseline_as_the_loss(
        tmp_path):
    ws = _book(tmp_path)["Stopped buying"]
    assert [v for v in _column(ws, "Agency / customer") if v] == ["Gone"]
    assert _column(ws, "Change (BDT)")[0] == -4_000_000


def test_an_empty_group_still_gets_a_sheet_that_says_so(tmp_path):
    res = sm.build([_row("Only", 0, 900_000)], AUG)
    ws = _book(tmp_path, res)["Declined"]
    assert "Nobody fell into this group." in _text(ws)


def test_a_thin_baseline_is_marked_rather_than_smoothed_away(tmp_path):
    rows = [_row("Patchy", 0, 200_000), _row("Patchy", 2, 3_000_000)]
    ws = _book(tmp_path, sm.build(rows, AUG))["Declined"]
    assert _column(ws, "Months traded")[0] == "1 of 3"


def test_a_warning_about_the_data_lands_at_the_top_of_the_summary(tmp_path):
    res = _result(data_last_day=date(2026, 8, 20))
    said = _text(_book(tmp_path, res)["Summary"])
    assert "READ THIS FIRST" in said
    assert "20 Aug 2026" in said


def test_the_summary_rolls_the_money_up_by_zone_and_by_sales_person(tmp_path):
    ws = _book(tmp_path)["Summary"]
    said = _text(ws)
    assert "BY ZONE" in said and "BY SALES PERSON" in said
    zones = [v for v in _column(ws, "By Zone", limit=8) if v]
    assert "Zone-9" in zones and "Zone-4" in zones
    reps = [v for v in _column(ws, "By Sales Person", limit=8) if v]
    assert "Rep A" in reps and "Rep B" in reps


def test_the_floor_is_reported_so_nobody_wonders_where_a_customer_went(
        tmp_path):
    rows = [_row("Tiny", 0, 100)] + [_row("Tiny", i, 900) for i in (1, 2, 3)]
    res = sm.build(rows, AUG)
    assert res.below_floor == 1
    assert "1 customer(s) fell below" in _text(_book(tmp_path, res)["Summary"])


def test_the_filename_names_the_period(tmp_path):
    assert smr.default_filename(AUG) == "Sales_Movement_Aug2026.xlsx"


# --------------------------------------------------------------------------
# the cells that were printing a number the caption says is never printed
# --------------------------------------------------------------------------
def test_the_stopped_buying_sheet_prints_no_percentage_at_all(tmp_path):
    """The caption on this sheet reads 'No percentage is printed'. It was
    writing -1 into every Change % cell, which Excel renders as -100%."""
    ws = _book(tmp_path)["Stopped buying"]
    assert [v for v in _column(ws, "Agency / customer") if v] == ["Gone"]
    assert _column(ws, "Change %")[0] is None


def test_a_refunder_gets_its_own_sheet_with_the_real_negative_figure(tmp_path):
    rows = [_row("Refunder", 0, -250_000)] + \
        [_row("Refunder", i, 4_000_000) for i in (1, 2, 3)]
    ws = _book(tmp_path, sm.build(rows, AUG))["Refunded more than sold"]
    assert [v for v in _column(ws, "Agency / customer") if v] == ["Refunder"]
    assert _column(ws, "This period (BDT)")[0] == -250_000
    assert _column(ws, "Change (BDT)")[0] == -4_250_000   # exceeds the baseline
    assert _column(ws, "Change %")[0] is None


def test_a_refunder_is_never_filed_under_stopped_buying(tmp_path):
    rows = [_row("Refunder", 0, -250_000)] + \
        [_row("Refunder", i, 4_000_000) for i in (1, 2, 3)]
    wb = _book(tmp_path, sm.build(rows, AUG))
    assert "Nobody fell into this group." in _text(wb["Stopped buying"])


# --------------------------------------------------------------------------
# a ticket count is not an amount of money
# --------------------------------------------------------------------------
def _tickets_settings():
    return sm.Settings(period_from=date(2026, 8, 1), period_to=date(2026, 8, 31),
                       trailing=3, threshold=0.20, floor=50,
                       measure=sm.MEASURE_TICKETS)


def test_the_tickets_measure_does_not_label_its_columns_bdt(tmp_path):
    """Choosing 'Tickets issued' used to present a count of tickets under a
    heading that said BDT, on every sheet in the workbook."""
    s = _tickets_settings()
    rows = [_row("Falling", 0, 300)] + \
        [_row("Falling", i, 1_000) for i in (1, 2, 3)]
    wb = _book(tmp_path, sm.build(rows, s))
    ws = wb["Declined"]
    assert _column(ws, "Baseline average (tickets)")[0] == 1_000
    assert _column(ws, "This period (tickets)")[0] == 300
    assert "(BDT)" not in _text(ws)


def test_the_tickets_summary_does_not_say_money(tmp_path):
    s = _tickets_settings()
    rows = [_row("Falling", 0, 300)] + \
        [_row("Falling", i, 1_000) for i in (1, 2, 3)]
    said = _text(_book(tmp_path, sm.build(rows, s))["Summary"])
    assert "LOST (TICKETS)" in said.upper()
    assert "MONEY LOST" not in said.upper()
    assert "50 tickets" in said            # the floor is stated in its unit


# --------------------------------------------------------------------------
# these sheets are built to be forwarded, and the names on them are typed by
# people -- the warehouse already holds one starting with a minus sign
# --------------------------------------------------------------------------
def test_a_name_starting_with_a_dash_survives_exactly_as_written(tmp_path):
    """In an .xlsx a cell carries its own type, and '-Ul-Islam Md Armaan' is
    written as typed text, which Excel displays and never evaluates. So it
    must come back byte for byte -- an apostrophe bolted on the front would
    be stored IN the string and read back by everyone."""
    rows = [_row("-Ul-Islam Md Armaan", 0, 100_000)] + \
        [_row("-Ul-Islam Md Armaan", i, 4_000_000) for i in (1, 2, 3)]
    ws = _book(tmp_path, sm.build(rows, AUG))["Declined"]
    assert _column(ws, "Agency / customer")[0] == "-Ul-Islam Md Armaan"


def test_a_name_starting_with_equals_is_text_not_a_live_formula(tmp_path):
    """'=' is the one lead character openpyxl turns into <f>."""
    rows = [_row("=SUM(A1:A9) Travels", 0, 100_000)] + \
        [_row("=SUM(A1:A9) Travels", i, 4_000_000) for i in (1, 2, 3)]
    ws = _book(tmp_path, sm.build(rows, AUG))["Declined"]
    cell = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").startswith("=SUM"):
            cell = ws.cell(row=r, column=1)
    assert cell is not None, "the name was not written at all"
    assert cell.value == "=SUM(A1:A9) Travels"      # intact, not mangled
    assert cell.data_type == "s"                    # text, not a formula


def test_an_ordinary_name_is_left_exactly_as_it_is(tmp_path):
    ws = _book(tmp_path)["Declined"]
    assert _column(ws, "Agency / customer")[0] == "Falling"


def test_a_phone_number_keeps_its_plus(tmp_path):
    """+880... is already safe as typed text. Mangling it would corrupt
    every number on the call list."""
    from openpyxl import Workbook

    from src.counter_master import _cell
    ws = Workbook().active
    _cell(ws, 1, 1, "+8801711000000")
    assert ws.cell(row=1, column=1).value == "+8801711000000"


def test_the_windows_table_does_not_say_averaged_when_nothing_is_averaged(
        tmp_path):
    """Last-year is a single window. Labelling it 'averaged into the
    baseline' describes an arithmetic that did not happen."""
    s = sm.Settings(period_from=date(2026, 8, 1), period_to=date(2026, 8, 31),
                    baseline=sm.BASELINE_LAST_YEAR)
    ws = _book(tmp_path, sm.build([], s))["Summary"]
    roles = [v for v in _column(ws, "Role", limit=4) if v]
    assert roles == ["measured", "the same period a year earlier"]


def test_the_windows_table_still_says_averaged_for_a_trailing_baseline(
        tmp_path):
    ws = _book(tmp_path)["Summary"]
    roles = [v for v in _column(ws, "Role", limit=6) if v]
    assert roles[1] == "averaged into the baseline"


def test_the_windows_table_prints_the_day_count_of_each_window(tmp_path):
    """A 19-day period against a 31-day one is the comparison that quietly
    invalidates itself, so the day counts are on the sheet."""
    ws = _book(tmp_path)["Summary"]
    days = [v for v in _column(ws, "Days", limit=6) if v]
    assert days[:2] == [31, 31]           # Aug 2026 and Jul 2026


# --------------------------------------------------------------------------
# a refusal is a page saying why, not a workbook of empty tabs
# --------------------------------------------------------------------------
def _refused():
    s = sm.Settings(period_from=date(2026, 2, 1), period_to=date(2026, 2, 28),
                    baseline=sm.BASELINE_LAST_YEAR)
    return sm.build([_row("Anyone", 0, 9_000_000)], s,
                    data_first_day=date(2025, 5, 1))


def test_a_refused_run_writes_only_a_page_saying_why(tmp_path):
    out = tmp_path / "refused.xlsx"
    smr.build_workbook(_refused(), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary"]          # no bucket tabs at all
    said = _text(wb["Summary"])
    assert "no data for the baseline" in said
    assert "none of it would have been true" in said


def test_a_refused_run_has_no_declined_tab_to_misread(tmp_path):
    """An empty 'Declined' beside a refusal reads as 'nobody declined'."""
    out = tmp_path / "refused2.xlsx"
    smr.build_workbook(_refused(), out)
    assert "Declined" not in load_workbook(out).sheetnames


def test_the_summary_states_what_the_measure_leaves_out(tmp_path):
    said = _text(_book(tmp_path)["Summary"])
    assert "Not counted here" in said
    assert "penalties and reissue adjustments" in said


# --------------------------------------------------------------------------
# both baselines on one sheet
# --------------------------------------------------------------------------
BOTH = sm.Settings(period_from=date(2026, 8, 1), period_to=date(2026, 8, 31),
                   baseline=sm.BASELINE_BOTH, trailing=3, threshold=0.20,
                   floor=500_000)


def _both_result():
    rows = []
    rows += [_row("BothDown", 0, 2_000_000)] + \
        [_row("BothDown", i, 9_000_000) for i in (1, 2, 3)] + \
        [_row("BothDown", 4, 8_000_000)]
    rows += [_row("BlipOnly", 0, 3_000_000)] + \
        [_row("BlipOnly", i, 9_000_000) for i in (1, 2, 3)] + \
        [_row("BlipOnly", 4, 3_000_000)]
    return sm.build(rows, BOTH)


def test_the_year_column_is_headed_with_the_year_it_holds(tmp_path):
    ws = _book(tmp_path, _both_result())["Declined"]
    assert _column(ws, "Aug 2025 (BDT)")[0] == 8_000_000


def test_the_sheet_carries_both_percentages_side_by_side(tmp_path):
    ws = _book(tmp_path, _both_result())["Declined"]
    assert _column(ws, "Change %")[0] is not None
    assert _column(ws, "vs last year %")[0] is not None


def test_the_agreement_column_separates_a_real_decline_from_a_blip(tmp_path):
    ws = _book(tmp_path, _both_result())["Declined"]
    names = _column(ws, "Agency / customer")
    agree = _column(ws, "Do the two agree?")
    got = dict(zip(names, agree))
    assert got["BothDown"] == "down on both"
    assert got["BlipOnly"] == "down vs recent only"


def test_the_summary_tallies_the_agreement(tmp_path):
    said = _text(_book(tmp_path, _both_result())["Summary"])
    assert "DO THE TWO COMPARISONS AGREE?" in said
    assert "down on both" in said
    assert "down vs recent only" in said


def test_a_trailing_only_run_has_no_year_columns(tmp_path):
    ws = _book(tmp_path)["Declined"]
    assert "vs last year %" not in _text(ws)
    assert "DO THE TWO COMPARISONS AGREE?" not in _text(_book(tmp_path)["Summary"])
