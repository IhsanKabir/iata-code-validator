"""The pricing lookup, against the stand-in Zenith of test_pricing_check.

Tax detail pages are built like the live one: every exemption box is
unticked in the markup and a script ticks the exempt ones afterwards.
"""
from datetime import date

import pytest

from src import pricing_check as pc
from src import pricing_lookup as pl
from src import pricing_parse as pp
from tests.test_pricing_check import TODAY, FakeZenith, same_both_seasons

SUMMER, WINTER, SUMMER27 = 34198, 34199, 42946


def tax_list_with_ids(rows):
    head = "".join(f"<th>{h}</th>" for h in pp._TAX_HEADER)
    body = ""
    for tax_id, name, code, value in rows:
        body += (f'<tr><td>Tax Name <a href="" onclick="return Affiche('
                 f'{tax_id});">{name}</a></td><td>Tax Code {code}</td>'
                 f"<td>Departure Airport Dhaka</td><td>Arrival Airport "
                 f"Kolkata</td><td>Value {value}</td><td>International "
                 f"&#xe5ca;</td><td>Domestic </td><td>History</td></tr>")
    return f"<table><tr>{head}</tr>{body}</table>"


def tax_detail(code, name, exempt_fb=(), exempt_pax=()):
    boxes = "".join(
        f'<!--<input type="checkbox" name="ClasseTarifaireExempted" '
        f'value="{i}" id="c{i}" />--><input type="checkbox" '
        f'name="ClasseTarifaireExempted" value="{i}" id="c{i}" />'
        f'<label for="c{i}">{label}</label>'
        for i, label in ((11, "Y - one way | DAC-CCU"),
                         (12, "Y - return | DAC-CCU")))
    boxes += "".join(
        f'<input type="checkbox" name="TypePassagerExempt" value="{i}" '
        f'id="p{i}" /><label for="p{i}">{label}</label>'
        for i, label in ((1, "Adult(s)"), (3, "Infant(s)")))
    script = "".join(
        f"set_chk(document.taxe.ClasseTarifaireExempted, '{i}');"
        for i in exempt_fb) + "".join(
        f"set_chk(document.taxe.TypePassagerExempt, '{i}');"
        for i in exempt_pax)
    return (f'<input type="text" name="CodeTaxe" value="{code}" />'
            f'<input type="text" name="Libelle" value="{name}" />{boxes}'
            f"<script>{script}</script>")


class LookupZenith(FakeZenith):
    def __init__(self):
        super().__init__()
        self.opened = []

    def fare_basis_list(self):
        rows = super().fare_basis_list()
        for r in rows:
            r.update({"Level": 8, "MinStay": 0, "MaxStay": 360,
                      "AllowOneWay": r["Code"].endswith("O")})
        return rows

    def tax_search(self, form):
        from src import pricing_client as cl
        cl.check_allowed("POST", "https://z.test" + cl.TAXES, form)
        if form["hAction"] == "Affiche_Result":
            self.opened.append((form["Id_TaxeAeropackHidden"],
                                form["RecEffDate"]))
            if form["Id_TaxeAeropackHidden"] == "1669":
                return tax_detail("BD", "Embarkation Fee", exempt_fb=[12],
                                  exempt_pax=[3])
            return tax_detail("YQ", "Carrier fee")
        return tax_list_with_ids([(1669, "Embarkation Fee", "BD", "500 BDT"),
                                  (4919, "Carrier fee", "YQ", "5 USD")])


def zenith():
    z = same_both_seasons(LookupZenith())
    for season in (SUMMER, WINTER):
        for row in z.fares[(11, season)]:
            row.update({"Modifier": "A. Staff",
                        "Modified": "/Date(1781085720000)/"})
    return z


def look(z, tmp_path, **kw):
    kw.setdefault("season_ids", [SUMMER, WINTER])
    return pl.run_lookup(z, pc.RouteIndex(tmp_path / "none.json"),
                         pl.LookupSettings(**kw), today=TODAY)


# ---- choosing what to read ---------------------------------------------------

def test_fare_basis_codes_and_patterns_are_matched():
    rows = LookupZenith().fare_basis_list()
    got, unmatched = pl.match_fare_bases("ybdccuo, YBD*R, NOPE", rows)
    assert sorted(r["Code"] for r in got) == ["YBDCCUO", "YBDCCUR"]
    assert unmatched == ["NOPE"]
    # a disabled fare basis is found too -- it is a read-out, not a check
    got, _ = pl.match_fare_bases("OLD*", rows)
    assert [r["Code"] for r in got] == ["OLDX"]


def test_a_season_already_over_is_read_on_a_day_inside_it():
    seasons = pp.parse_fare_editor_page(
        __import__("tests.test_pricing_check", fromlist=["fe_page"]).fe_page()
    ).seasons
    summer = seasons[0]                      # 28 Mar - 24 Oct 2026
    assert pl.lookup_date(summer, date(2026, 12, 1)) == date(2026, 4, 4)
    assert pl.lookup_date(summer, TODAY) == pc.check_date(summer, TODAY)


def test_the_tax_detail_exemptions_are_read_from_the_page_script():
    d = pp.parse_tax_detail(tax_detail("BD", "Embarkation Fee",
                                       exempt_fb=[12], exempt_pax=[3]))
    assert (d.code, d.name) == ("BD", "Embarkation Fee")
    assert d.exempt_fare_bases() == {12}
    assert d.exempt["TypePassagerExempt"] == {"3": "Infant(s)"}
    assert pp.parse_tax_detail(tax_detail("YQ", "x")).exempt == {}
    with pytest.raises(pp.LayoutChanged):
        pp.parse_tax_detail("<html>no lists</html>")


# ---- the read-out ------------------------------------------------------------

def test_a_fare_basis_is_read_out_in_full_for_each_season(tmp_path):
    res = look(zenith(), tmp_path, fare_bases_text="YBDCCUO")
    assert res.routes == ["DAC-CCU"]
    assert res.fare_bases[0]["Code"] == "YBDCCUO"
    assert res.fare_bases[0]["MaxStay"] == 360
    adult = [r for r in res.fares if "Adult" in str(r["Passenger"])]
    assert [r["Season"] for r in adult] == ["Summer 2026", "Winter 2026-27"]
    assert adult[0]["AmountOW"] == 300 and adult[0]["Currency"] == "USD"
    assert adult[0]["Changed by"] == "A. Staff"
    assert adult[0]["Changed on"].startswith("2026-06-1")


def test_taxes_show_who_is_exempt_and_each_tax_opens_once(tmp_path):
    z = zenith()
    res = look(z, tmp_path, fare_bases_text="YBDCCU*")
    bd = [t for t in res.taxes if t["Tax code"] == "BD"]
    assert [t["Season"] for t in bd] == ["Summer 2026", "Winter 2026-27"]
    assert bd[0]["Looked-up fare bases exempt"] == "YBDCCUR"
    assert bd[0]["Exempt: Passenger types"] == "Infant(s)"
    yq = next(t for t in res.taxes if t["Tax code"] == "YQ")
    assert yq["Looked-up fare bases exempt"] == "none"
    # two taxes x two season days, each opened once
    assert len(z.opened) == len(set(z.opened)) == 4


def test_exemptions_can_be_left_out_to_save_time(tmp_path):
    z = zenith()
    res = look(z, tmp_path, fare_bases_text="YBDCCUO", tax_exemptions=False)
    assert res.taxes and not z.opened
    assert "Looked-up fare bases exempt" not in res.taxes[0]


def test_penalties_and_quick_sale_are_read_for_the_fare_basis(tmp_path):
    res = look(zenith(), tmp_path, fare_bases_text="YBDCCUO")
    assert {p["Season"] for p in res.penalties} == {"Summer 2026",
                                                    "Winter 2026-27"}
    assert res.penalties[0]["Change fee"] == "300 USD"
    # the result says "Y one way" (the description); the code is found
    assert res.penalties[0]["Fare basis"] == "YBDCCUO"
    assert res.penalties[0]["Fare class (as shown)"] == "Y one way"
    qs = [q for q in res.quick_sale if q["Fare basis"] == "YBDCCUO"]
    assert [q["OW"] for q in qs] == [37097, 37097]
    assert qs[0]["OW tax"] == 43100 - 37097


def test_a_fare_basis_quick_sale_does_not_list_is_said_so(tmp_path):
    z = zenith()
    z.qs[("DACCCU", "WINTER 2026-27", False, "BDT")] = [("YBDCCUR", None,
                                                         61828)]
    res = look(z, tmp_path, fare_bases_text="YBDCCUO")
    winter = [q for q in res.quick_sale if q["Season"] == "Winter 2026-27"]
    assert winter == [{"Season": "Winter 2026-27", "Route": "DAC-CCU",
                       "Travel date": "31OCT26", "Fare basis": "YBDCCUO",
                       "Currency": "not listed in Quick Sale"}]


def test_a_route_alone_needs_the_route_list(tmp_path):
    with pytest.raises(ValueError, match="route list"):
        look(zenith(), tmp_path, routes_text="DAC-CCU")
    idx = pc.RouteIndex(tmp_path / "routes.json")
    idx.legs = {11: {"DAC-CCU"}, 12: {"DAC-CCU"}}
    idx.built = "2026-10-05 16:58"
    idx.save()
    res = pl.run_lookup(zenith(), idx, pl.LookupSettings(
        routes_text="DAC-CCU", include_return=False,
        season_ids=[SUMMER]), today=TODAY)
    assert [f["Code"] for f in res.fare_bases] == ["YBDCCUO", "YBDCCUR"]
    assert any("route list built 2026-10-05" in n for n in res.notes)


def test_too_broad_a_lookup_asks_to_narrow_it(tmp_path, monkeypatch):
    monkeypatch.setattr(pl, "MAX_FARE_BASES", 1)
    with pytest.raises(ValueError, match="narrow"):
        look(zenith(), tmp_path, fare_bases_text="YBD*")


def test_nothing_to_look_up_is_refused(tmp_path):
    with pytest.raises(ValueError, match="fare basis or a route"):
        look(zenith(), tmp_path)
    with pytest.raises(ValueError, match="season"):
        look(zenith(), tmp_path, fare_bases_text="YBDCCUO", season_ids=[])


def test_the_lookup_workbook_has_a_sheet_per_part(tmp_path):
    from openpyxl import load_workbook
    from src import pricing_lookup_report as plr
    res = look(zenith(), tmp_path, fare_bases_text="YBDCCUO")
    path = tmp_path / "lookup.xlsx"
    plr.build_workbook(res, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["About", "Fare basis", "Fares", "Penalties",
                             "Taxes", "Quick Sale"]
    heads = [c.value for c in wb["Fares"][4]]
    assert heads[:4] == ["Season", "Route", "Fare basis", "Passenger"]
    assert "OW" in heads and "Changed by" in heads
    assert "Miles OW" not in heads          # empty on every row: left out
    taxes = [c.value for c in wb["Taxes"][4]]
    assert "Looked-up fare bases exempt" in taxes
