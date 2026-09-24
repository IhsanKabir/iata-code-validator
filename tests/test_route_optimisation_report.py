"""The Route Optimisation workbook."""
from datetime import date

from openpyxl import load_workbook

from src import route_optimisation as ro
from src import route_optimisation_report as ror
from tests.test_route_optimisation import _legs, _market


def _res():
    return ro.build(_legs(), _market(), days_observed=14,
                    distances={"DAC-CCU": 329.656},
                    revenue={"DAC-CCU": {"2026-09": 10_440_000}})


def _book(tmp_path):
    out = tmp_path / "ro.xlsx"
    ror.build_workbook(_res(), out)
    return load_workbook(out)


def _text(ws) -> str:
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row
                    if c.value is not None)


def _column(ws, header, limit=40):
    for r in range(1, ws.max_row + 1):
        for c in range(1, 22):
            if ws.cell(row=r, column=c).value == header:
                return [ws.cell(row=i, column=c).value
                        for i in range(r + 1, min(r + 1 + limit,
                                                  ws.max_row + 1))]
    raise AssertionError(f"no column headed {header!r}")


def test_the_workbook_has_the_three_sheets(tmp_path):
    assert _book(tmp_path).sheetnames == ["Summary", "Every route",
                                          "Who flies what"]


def test_the_summary_leads_with_the_routes_worth_acting_on(tmp_path):
    said = _text(_book(tmp_path)["Summary"])
    assert "FULL AND OUT-FLOWN" in said
    assert "DAC-CCU" in said


def test_the_summary_says_seats_decide_it_not_departures(tmp_path):
    said = _text(_book(tmp_path)["Summary"])
    assert "180 seats to 72" in said


def test_both_shares_are_shown_because_they_disagree(tmp_path):
    ws = _book(tmp_path)["Every route"]
    seat = _column(ws, "Seat share")[0]
    flight = _column(ws, "Flight share")[0]
    assert seat < flight        # an ATR against an A320


def test_the_caveats_are_on_the_sheet_not_in_a_docstring(tmp_path):
    said = _text(_book(tmp_path)["Summary"])
    assert "READ THIS FIRST" in said
    assert "is a floor" in said


def test_every_rival_says_where_its_seat_count_came_from(tmp_path):
    ws = _book(tmp_path)["Who flies what"]
    assert set(_column(ws, "Seat count from")[:2]) <= {"airline", "type",
                                                       "operator"}


def test_a_thin_route_is_kept_and_labelled_rather_than_ranked(tmp_path):
    thin = ro.build(_legs()[:1], _market(), days_observed=14)
    out = tmp_path / "thin.xlsx"
    ror.build_workbook(thin, out)
    ws = load_workbook(out)["Every route"]
    assert "too few flights to judge" in _text(ws)


def test_a_run_with_nothing_squeezed_says_so(tmp_path):
    from tests.test_route_optimisation import _Leg, _sched
    legs = [_Leg("DAC-CCU", 189, 90, date(2026, 9, 1 + i), "Boeing 737-800")
            for i in range(14)]
    rows = [_sched("BG", "9", f"2026-09-{i+1:02d}T08:00", "DH8", 30)
            for i in range(14)]
    out = tmp_path / "none.xlsx"
    ror.build_workbook(ro.build(legs, rows, days_observed=14), out)
    assert "No route is both full and out-flown." in \
        _text(load_workbook(out)["Summary"])


# ---- RASK is a live formula, so a typed-in distance computes it -----------

def _no_distance_pair():
    """Both directions full and out-flown, and missing from the table."""
    def view(route):
        return ro.RouteView(route, our_seats=72, our_taken=70, our_legs=20,
                            rivals=[ro.Rival("6E", 2, 360)],
                            rask_revenue=1_000_000.0, rask_seats=2_000.0,
                            rask_months=("2026-07",))
    return ro.Result(routes=[view("DAC-XYZ"), view("XYZ-DAC")])


def _rows_for(ws, label):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == label:
            return r
    raise AssertionError(f"no row {label!r}")


def test_a_missing_distance_is_left_blank_with_rask_as_a_formula(tmp_path):
    out = tmp_path / "ro.xlsx"
    ror.build_workbook(_no_distance_pair(), out)
    ws = load_workbook(out)["Every route"]
    r = _rows_for(ws, "→ DAC-XYZ")
    assert ws[f"J{r}"].value is None                     # blank, to be typed
    assert ws[f"P{r}"].value == 1_000_000 and ws[f"Q{r}"].value == 2_000
    rask = ws[f"L{r}"].value
    assert rask.startswith("=") and f"J{r}" in rask and f"P{r}" in rask


def test_one_typed_distance_fills_the_pair(tmp_path):
    """The inbound reads the outbound's distance; the pair row pools both."""
    out = tmp_path / "ro.xlsx"
    ror.build_workbook(_no_distance_pair(), out)
    ws = load_workbook(out)["Every route"]
    o = _rows_for(ws, "→ DAC-XYZ")
    i = _rows_for(ws, "← XYZ-DAC")
    assert f"J{o}" in ws[f"J{i}"].value
    pair = _rows_for(ws, "DAC ⇄ XYZ")
    assert all(f"{c}{x}" in ws[f"L{pair}"].value
               for c in "JPQ" for x in (o, i))


def test_the_summary_reads_its_distances_from_every_route(tmp_path):
    out = tmp_path / "ro.xlsx"
    ror.build_workbook(_no_distance_pair(), out)
    wb = load_workbook(out)
    o = _rows_for(wb["Every route"], "→ DAC-XYZ")
    s = _rows_for(wb["Summary"], "→ DAC-XYZ")
    assert f"'Every route'!J{o}" in wb["Summary"][f"J{s}"].value


def test_a_known_distance_is_written_as_a_number(tmp_path):
    ws = _book(tmp_path)["Every route"]
    r = _rows_for(ws, "→ DAC-CCU")
    assert ws[f"J{r}"].value == 329.656
