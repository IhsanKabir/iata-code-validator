"""The FFP account search and the collection built on it.

Pages are built in the structure captured from Zenith's
CustomerAccount/IframeSearch, with invented members -- no real customer
appears here.
"""
import json
import threading

import pytest
from openpyxl import load_workbook

from src import ffp_collect as fc
from src import zenith_ffp as zf


def _row(num, level="Silver", miles="0", last="Testsurname",
         first="Testfirst"):
    return (
        '<tr class="customerRow">'
        f'<input class="idCustomer" id="idCustomer" name="idCustomer" '
        f'type="hidden" value="{num}" />'
        f'<input class="customerFullName" id="customerFullName" '
        f'name="customerFullName" type="hidden" value="{first} {last}" />'
        '<td> <span class="label label-default"> <a style="color: white" '
        f'href="/x/TravelAgency.ashx?IdCustomer={num}" target="_blank">{num}'
        '</a> <span class="glyphicon glyphicon-comment" /> </span> </td>'
        f'<td><span class="surname">{last}</span></td>'
        f'<td><span class="firstname">{first} </span></td>'
        '<td> <span class="birthdate">01/01/1990</span> </td>'
        '<td style="width: 125px"><span class="email">a@example.test</span></td>'
        '<td><span class="phone">+8800000000000</span></td>'
        '<td><span class="documentId">X0000000</span></td>'
        f'<td><span>{level}</span></td>'
        f'<td><span>{miles}</span></td>'
        '<td><button class="selectFFPCustomer">Select</button></td>'
        '</tr>')


def _page(rows, total=None, token="tok"):
    total = len(rows) if total is None else total
    note = ("<div> Results list not displaying all results. Please restrict "
            "your search criterias. </div>") if total > len(rows) else ""
    return (
        '<html><body><form action="/Zenith/x/CustomerAccount/IframeSearch" '
        'id="formIframeSearch" method="post">'
        f'<input name="__RequestVerificationToken" type="hidden" '
        f'value="{token}" /></form>'
        '<li class="active"><a href="#ffpSearchResults">FFP Results '
        f'<span class="badge">{total}</span></a></li>{note}'
        '<table id="results"><tbody>' + "".join(rows) +
        "</tbody></table></body></html>")


# ---- reading the page -------------------------------------------------------

def test_a_member_row_is_read_column_by_column():
    page = zf.parse_search_results(_page([_row("12382185", "Gold", "70")]))
    m = page.members[0]
    assert (m.ffp_number, m.customer_id, m.level, m.miles) == \
        ("12382185", "12382185", "Gold", "70")
    assert (m.last_name, m.first_name) == ("Testsurname", "Testfirst")
    assert m.birth_date == "01/01/1990" and m.id_number == "X0000000"


def test_the_badge_is_the_full_count_even_when_only_50_are_listed():
    """As captured: 'FFP Results 119847', 50 rows, and a warning."""
    rows = [_row(str(10000000 + i)) for i in range(50)]
    page = zf.parse_search_results(_page(rows, total=119847))
    assert page.total == 119847 and len(page.members) == 50
    assert page.truncated and not page.complete


def test_a_small_search_is_complete():
    page = zf.parse_search_results(_page([_row("1"), _row("2")]))
    assert page.complete and page.total == 2


def test_an_error_page_is_not_mistaken_for_no_results():
    with pytest.raises(zf.FFPSearchError):
        zf.parse_search_results("<html>504 Gateway Timeout</html>")


def test_the_security_token_is_read():
    assert zf.parse_token(_page([])) == "tok"
    with pytest.raises(zf.FFPSearchError):
        zf.parse_token("<html></html>")


# ---- collecting -------------------------------------------------------------

def _match(value, q, how):
    """How a fake field matches: 'prefix', 'contains', 'exact', or the
    same three needing a trailing '%' before they match partly."""
    if how.endswith("%"):
        if not q.endswith("%"):
            return value == q
        q, how = q[:-1], how[:-1]
    return {"prefix": value.startswith(q), "contains": q in value,
            "exact": value == q}[how]


class _FakeZenith:
    """An FFP search over invented members, with Zenith's 50-row limit.

    `number` / `name` say how the FFP number and last name fields match
    (see _match); `case` whether last names are case-sensitive.
    """

    def __init__(self, members, number="prefix", name="exact", case=False,
                 honours=None, server_cap=None):
        self.members = members
        self.number, self.name, self.case = number, name, case
        # ("form" | "query", field name): a larger page it will list
        self.honours, self.server_cap = honours, server_cap
        self.searches = 0
        self.calls = []

    def _cap(self, extra_form, extra_query):
        if self.honours:
            where, name = self.honours
            got = (extra_form if where == "form" else extra_query) or {}
            if name in got:
                asked = int(got[name])
                return min(asked, self.server_cap or asked)
        return 50

    def search(self, *, level="", ffp_number="", last_name="",
               extra_form=None, extra_query=None, **_):
        self.searches += 1
        self.calls.append((level, ffp_number, last_name))
        last_name = last_name.rstrip()      # as Zenith: trailing space trimmed
        norm = (lambda v: v) if self.case else (lambda v: v.upper())
        hits = [m for m in self.members
                if (not level or m.level == level)
                and (not ffp_number
                     or _match(m.ffp_number, ffp_number, self.number))
                and (not last_name
                     or _match(norm(m.last_name), norm(last_name),
                               self.name))]
        cap = self._cap(extra_form, extra_query)
        return zf.SearchPage(total=len(hits), members=tuple(hits[:cap]),
                             truncated=len(hits) > cap)


_SYLLABLES = ["Ah", "med", "Kha", "n", "Ra", "hman", "Is", "lam", "Ho",
              "ssain", "Chow", "dhury", "Sar", "kar", "Al Am", "in"]


def _members(n, start=10_000_000, step=7):
    levels = ["Silver"] * 17 + ["Gold"] * 2 + ["Platinum"]
    out = []
    for i in range(n):
        a = _SYLLABLES[(i * 3) % len(_SYLLABLES)]
        b = _SYLLABLES[(i * 5 + 1) % len(_SYLLABLES)].lower()
        last = f"{a}{b}{chr(97 + i % 26)}{chr(97 + (i // 26) % 26)}"
        out.append(zf.FFPMember(str(start + i * step), str(start + i * step),
                                last, "F", "", "", "", "", levels[i % 20],
                                "0"))
    return out


@pytest.fixture
def store(tmp_path):
    s = fc.FFPStore(tmp_path / "ffp.sqlite")
    yield s
    s.close()


@pytest.mark.parametrize("mode", ["prefix", "contains"])
def test_every_member_is_collected_past_the_50_row_limit(store, mode):
    people = _members(1_234)
    fake = _FakeZenith(people, number=mode)
    got = fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert got.members == len(people)
    assert store.member_count() == 1_234
    assert store.get("finished") == "1"


def test_the_level_counts_are_kept_to_check_the_list_against(store):
    people = _members(400)
    fc.collect(_FakeZenith(people), store, stop=threading.Event(), delay_s=0)
    totals = store.level_totals()
    assert totals["Silver"] == sum(m.level == "Silver" for m in people)
    assert store.counts_by_level()["Silver"] == totals["Silver"]
    assert totals["Titanium"] == 0


def test_nothing_partial_stops_the_run_instead_of_a_partial_list(store):
    with pytest.raises(fc.PartialMatchUnsupported):
        fc.collect(_FakeZenith(_members(300), number="exact", name="exact"),
                   store, stop=threading.Event(), delay_s=0)
    assert store.get("finished") != "1"
    # the level counts are kept, so the tab can still show them
    assert store.level_totals()["Silver"] > 0


def test_a_number_wildcard_is_found_and_used(store):
    """As on the live system, a plain partial number finds nothing."""
    people = _members(700)
    fake = _FakeZenith(people, number="prefix%")
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.get("mode") == "number" and store.get("wild") == "%"
    assert store.member_count() == 700


@pytest.mark.parametrize("case", [False, True])
def test_whole_numbers_only_falls_back_to_last_names(store, case):
    people = _members(1_500)
    fake = _FakeZenith(people, number="exact", name="prefix", case=case)
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.get("mode") == "name"
    assert store.get("case") == ("1" if case else "0")
    assert store.member_count() == 1_500
    assert store.get("finished") == "1"


def test_a_last_name_wildcard_is_found_and_used(store):
    people = _members(600)
    fake = _FakeZenith(people, number="exact", name="contains%")
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert (store.get("mode"), store.get("wild")) == ("name", "%")
    assert store.member_count() == 600


def test_stop_keeps_progress_and_the_next_run_carries_on(store):
    people = _members(900)
    fake = _FakeZenith(people)
    stop = threading.Event()
    seen = []

    def halt_after_five(p):
        seen.append(p)
        if len(seen) == 5:
            stop.set()

    fc.collect(fake, store, stop=stop, on_progress=halt_after_five, delay_s=0)
    assert store.get("finished") == "0"
    finished = {p for p, in store.con.execute(
        "SELECT prefix FROM frontier WHERE state != 'pending'")}
    assert finished

    again = _FakeZenith(people)
    fc.collect(again, store, stop=threading.Event(), delay_s=0)
    assert store.member_count() == 900
    # the method check is not repeated -- only a recount of each level
    # still being collected -- and finished searches are not redone
    recounts = [lv for lv, num, name in again.calls if lv and not name]
    assert len(recounts) == len(set(recounts)) <= 4
    assert not finished & {num for _, num, _ in again.calls if num}


def test_the_workbook_lists_every_member_and_checks_the_counts(store,
                                                               tmp_path):
    people = _members(120)
    fc.collect(_FakeZenith(people), store, stop=threading.Event(), delay_s=0)
    out = tmp_path / "ffp.xlsx"
    assert fc.export_workbook(store, out) == 120
    wb = load_workbook(out, read_only=True)
    assert wb.sheetnames == ["Summary", "Members"]
    rows = list(wb["Members"].iter_rows(values_only=True))
    assert rows[0][0] == "FFP number" and len(rows) == 121
    summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r]
    assert "CONFIDENTIAL" in summary[0][0]
    silver = next(r for r in summary if r[0] == "Silver")
    assert silver[1] == silver[2] and silver[3] == 0
    assert ("Collection", "complete") in [tuple(r[:2]) for r in summary]


def test_start_over_clears_everything(store):
    fc.collect(_FakeZenith(_members(60)), store, stop=threading.Event(),
               delay_s=0)
    store.reset()
    assert store.member_count() == 0 and store.get("mode") == ""


# ---- talking to Zenith ------------------------------------------------------

class _Resp:
    def __init__(self, text, status=200, url="https://z.test/x"):
        self.text, self.status_code, self.url = text, status, url


class _HttpStub:
    """Stands in for requests.Session: GETs answer PollSession and the
    empty form; POSTs pop the next scripted response."""

    def __init__(self, posts, poll_ok=True):
        self.posts = list(posts)
        self.poll_ok = poll_ok
        self.sent = []
        self.gets = []

    def get(self, url, **_):
        self.gets.append(url)
        if "PollSession" in url:
            return _Resp("PollSession Successful." if self.poll_ok else "no")
        return _Resp(_page([], token="form-token"))

    def post(self, url, data=None, **_):
        self.sent.append(dict(data))
        return self.posts.pop(0)


class _Session:
    def __init__(self, http):
        self.session = http
        self.state_values = {"ID_ADMIN": "1", "ID_SOCIETE": "2"}


def _searcher(posts, **kw):
    http = _HttpStub(posts, **kw)
    return zf.FFPSearcher(_Session(http), base_url="https://z.test"), http


def test_a_search_opens_the_modern_session_and_sends_the_form_token(
        monkeypatch):
    s, http = _searcher([_Resp(_page([_row("1")], token="next"))])
    page = s.search(level="Gold")
    assert page.total == 1
    assert any("PollSession?idUser=1&idCompany=2" in u for u in http.gets)
    assert http.sent[0]["__RequestVerificationToken"] == "form-token"
    assert http.sent[0]["LevelName"] == "Gold"


def test_a_timeout_is_retried(monkeypatch):
    monkeypatch.setattr(zf.time, "sleep", lambda _s: None)
    s, http = _searcher([_Resp("gateway", status=504),
                         _Resp(_page([_row("1")]))])
    assert s.search(ffp_number="1").total == 1 and len(http.sent) == 2


def test_a_lost_modern_session_is_reopened_once(monkeypatch):
    s, http = _searcher([_Resp("<html>ERROR PAGE</html>"),
                         _Resp(_page([_row("1")]))])
    assert s.search(ffp_number="1").total == 1
    assert sum("PollSession" in u for u in http.gets) == 2


def test_a_session_zenith_refuses_asks_for_a_new_sign_in():
    s, _ = _searcher([], poll_ok=False)
    with pytest.raises(zf.FFPSessionError):
        s.search(level="Gold")


# ---- order, level choice, several searches at once --------------------------

def _by_name(n=900):
    return _FakeZenith(_members(n), number="exact", name="prefix")


def _after_method_check(fake, store):
    """Run the method check alone; return where collection searches start,
    so its probe searches are not counted as collection."""
    fc._choose_method(fake, store, lambda _m: None)
    return len(fake.calls)


def test_small_levels_are_collected_before_silver(store):
    fake = _by_name()
    start = _after_method_check(fake, store)
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    order = [lv for lv, num, name in fake.calls[start:] if name]
    firsts = {lv: order.index(lv) for lv in ("Gold", "Platinum", "Silver")
              if lv in order}
    lasts = {lv: len(order) - 1 - order[::-1].index(lv) for lv in firsts}
    # a level of 50 or fewer is complete from its count and needs none
    assert "Gold" in firsts and "Silver" in firsts
    for level in firsts:
        if level != "Silver":
            assert lasts[level] < firsts["Silver"]


def test_a_run_can_leave_silver_for_later(store):
    people = _members(1_200)
    fake = _FakeZenith(people, number="exact", name="prefix")
    fc.collect(fake, store, stop=threading.Event(), delay_s=0,
               levels=["Gold", "Platinum", "Titanium"])
    found = store.counts_by_level()
    want = {lv: sum(m.level == lv for m in people) for lv in ("Gold",
                                                            "Platinum")}
    assert found["Gold"] == want["Gold"]
    assert found["Platinum"] == want["Platinum"]
    assert fc.pending_levels(store) == ["Silver"]
    assert store.get("finished") == "0"


@pytest.mark.parametrize("workers", [2, 4])
def test_several_searches_at_once_still_find_everyone_once(store, workers):
    people = _members(1_100)
    fake = _FakeZenith(people, number="exact", name="prefix")
    start = _after_method_check(fake, store)
    fc.collect(fake, store, stop=threading.Event(), delay_s=0,
               workers=workers)
    assert store.member_count() == 1_100
    assert store.get("finished") == "1"
    names = [(lv, name) for lv, num, name in fake.calls[start:]]
    assert len(names) == len(set(names))          # no search run twice


def _flaky(fake, fail_on, exc):
    real = fake.search
    calls = {"n": 0}

    def search(**kw):
        calls["n"] += 1
        if calls["n"] in fail_on:
            raise exc
        return real(**kw)

    fake.search = search
    return real


def test_one_failed_search_does_not_end_the_run(store, monkeypatch):
    monkeypatch.setattr(fc, "RETRY_WAIT_S", 0)
    fake = _by_name()
    _flaky(fake, {30}, zf.FFPSearchError("Zenith returned 504 four times."))
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.member_count() == 900 and store.get("finished") == "1"


def test_five_failures_in_a_row_end_it_with_the_search_requeued(
        store, monkeypatch):
    monkeypatch.setattr(fc, "RETRY_WAIT_S", 0)
    fake = _by_name()
    real = _flaky(fake, set(range(30, 40)),
                  zf.FFPSearchError("Zenith returned 504 four times."))
    with pytest.raises(zf.FFPSearchError):
        fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.running() == 0                   # nothing left half-taken
    fake.search = real
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.member_count() == 900


def test_a_lost_sign_in_ends_the_run_at_once(store, monkeypatch):
    monkeypatch.setattr(fc, "RETRY_WAIT_S", 0)
    fake = _by_name()
    _flaky(fake, {30}, zf.FFPSessionError("signed out"))
    with pytest.raises(zf.FFPSessionError):
        fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.running() == 0


def test_no_two_punctuation_marks_are_added_in_a_row(store):
    store.put("mode", "name")
    kids = fc._children(store, "Gold|Md.")
    assert "Gold|Md.a" in kids
    assert not [k for k in kids if k[-2:] in ("..", ".-", ".'", ". ")]
    assert "Gold|Md a" in fc._children(store, "Gold|Md")


def test_names_starting_with_a_digit_are_searched(store):
    fake = _by_name(600)
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    searched = {name for _, _, name in fake.calls}
    assert {"0", "9"} <= searched


def test_a_saved_run_gets_digit_starts_once(store):
    store.put("mode", "name")
    store.put("space_fix", "1")
    store.put("level_total:Gold", 1050)
    store.put("level_total:Platinum", 20)
    fc.add_digit_starts(store)
    queued = {p for p, in store.con.execute("SELECT prefix FROM frontier")}
    assert "Gold|0" in queued and "Gold|9" in queued
    assert "Platinum|0" not in queued             # fits one search already
    assert store.get("digit_fix") == "1"


def test_the_limit_test_skips_a_setting_zenith_errors_on(store):
    fake = _FakeZenith(_members(2_000), name="prefix",
                       honours=("form", "Take"))
    real = fake.search

    def choke_on_pagesize(**kw):
        if "PageSize" in (kw.get("extra_form") or {}):
            raise zf.FFPSearchError("500 on an unknown field")
        return real(**kw)

    fake.search = choke_on_pagesize
    found = fc.probe_page_limit(fake, store, stop=threading.Event())
    assert found["name"] == "Take"


def test_a_row_zenith_listed_but_that_could_not_be_read_is_not_split(store):
    """Counting rows would call this incomplete for ever and split it."""
    page = zf.SearchPage(total=3, members=tuple(_members(2)), truncated=False)
    assert page.complete and page.short == 1


def test_a_row_without_a_readable_number_keeps_its_customer_id():
    row = _row("12345678").replace(">12345678</a>", ">—</a>")
    m = zf.parse_member(row)
    assert m.ffp_number == "12345678" == m.customer_id


def test_a_progress_file_from_the_last_version_still_opens(tmp_path):
    import sqlite3
    path = tmp_path / "old.sqlite"
    con = sqlite3.connect(path)
    con.executescript("""
        CREATE TABLE frontier (prefix TEXT PRIMARY KEY, state TEXT NOT NULL,
                               total INTEGER NOT NULL DEFAULT -1);
        INSERT INTO frontier VALUES ('Silver|A', 'pending', -1),
                                    ('Gold|B', 'pending', -1);""")
    con.commit()
    con.close()
    s = fc.FFPStore(path)
    assert s.next_pending() == "Gold|B"           # Gold outranks Silver
    s.close()


# ---- asking Zenith for more than 50 ----------------------------------------

@pytest.mark.parametrize("where,name", [("form", "MaxResults"),
                                        ("query", "NbReponse")])
def test_a_setting_zenith_honours_is_found(store, where, name):
    fake = _FakeZenith(_members(2_000), name="prefix", honours=(where, name))
    found = fc.probe_page_limit(fake, store, stop=threading.Event())
    assert (found["where"], found["name"]) == (where, name)
    assert found["listed"] > 50
    assert json.loads(store.get("page_param"))["name"] == name


def test_no_honoured_setting_means_50_is_a_hard_limit(store):
    fake = _FakeZenith(_members(2_000), name="prefix")
    assert fc.probe_page_limit(fake, store, stop=threading.Event()) is None
    assert store.get("page_param") == ""
    # every candidate was tried, in the form and in the address
    assert fake.searches == 1 + 2 * len(fc.LIMIT_NAMES)


def test_a_larger_page_collects_everything_in_far_fewer_searches(tmp_path):
    people = _members(3_000)
    plain = fc.FFPStore(tmp_path / "plain.sqlite")
    slow = _FakeZenith(people, number="exact", name="prefix")
    fc.collect(slow, plain, stop=threading.Event(), delay_s=0)

    big = fc.FFPStore(tmp_path / "big.sqlite")
    fast = _FakeZenith(people, number="exact", name="prefix",
                       honours=("form", "PageSize"), server_cap=500)
    fc.probe_page_limit(fast, big, stop=threading.Event(), level="Silver")
    fc.collect(fast, big, stop=threading.Event(), delay_s=0)

    assert big.member_count() == plain.member_count() == 3_000
    assert big.get("finished") == "1"
    assert fast.searches < slow.searches / 3
    plain.close()
    big.close()


def test_a_level_that_fits_one_page_drops_its_queued_searches(store):
    """Gold was part-collected 50 at a time; once a page holds all of it,
    one search finishes Gold and its queue is closed."""
    people = _members(2_000)
    slow = _FakeZenith(people, number="exact", name="prefix")
    stop = threading.Event()
    fc.collect(slow, store, stop=stop, delay_s=0,
               on_progress=lambda p: stop.set())   # stop after one search
    assert fc.pending_levels(store)
    fast = _FakeZenith(people, number="exact", name="prefix",
                       honours=("form", "PageSize"))
    fc.probe_page_limit(fast, store, stop=threading.Event(), level="Silver")
    before = fast.searches
    fc.collect(fast, store, stop=threading.Event(), delay_s=0)
    assert store.member_count() == 2_000
    assert fast.searches - before <= 4             # one per level, at most


# ---- Zenith trims a trailing space ------------------------------------------

def test_no_search_ends_in_a_space_and_two_word_names_are_found(store):
    """'A ' matched exactly what 'A' did on the live system, so splitting
    it again only added spaces. A space now always comes with a letter."""
    people = _members(1_500)           # includes "Al Am…" two-word names
    assert any(" " in m.last_name for m in people)
    fake = _FakeZenith(people, number="exact", name="contains")
    fc.collect(fake, store, stop=threading.Event(), delay_s=0)
    assert store.member_count() == 1_500
    assert not any(name.endswith(" ") for _, _, name in fake.calls if name)


def test_a_run_left_with_space_chains_is_repaired(store):
    store.put("mode", "name")
    store.seed(["Gold|A", "Gold|A ", "Gold|A  ", "Gold|Al  x", "Gold|B"])
    store.mark("Gold|A", fc.SPLIT, 219)
    store.mark("Gold|A ", fc.SPLIT, 219)
    dropped = fc.repair_space_chains(store)
    queued = {p for p, in store.con.execute(
        "SELECT prefix FROM frontier WHERE state='pending'")}
    assert dropped == 2                     # "A  " and "Al  x"
    assert "Gold|A a" in queued and "Gold|A z" in queued
    assert not any(p.endswith(" ") for p in queued)
    assert "Gold|B" in queued
    assert store.get("space_fix") == "1"


def test_each_run_recounts_the_levels_it_collects(store):
    people = _members(900)
    fc.collect(_by_name(900), store, stop=threading.Event(), delay_s=0,
               levels=["Gold"])
    store.put("level_total:Gold", 1)              # a stale count
    grown = _FakeZenith(people + _members(40, start=20_000_000),
                        number="exact", name="prefix")
    store.seed(["Gold|Zz"])                       # so Gold has work left
    fc.collect(grown, store, stop=threading.Event(), delay_s=0,
               levels=["Gold"])
    assert store.level_totals()["Gold"] == sum(
        m.level == "Gold" for m in grown.members)
