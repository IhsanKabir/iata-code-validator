"""The blank forms that make the missing columns exist."""
from datetime import date

import pytest
from openpyxl import load_workbook

from src import counter_master as cm
from src import report_templates as rt
from src import visit_master as vm


def _labels(ws):
    return [str(c.value or "") for row in ws.iter_rows() for c in row if c.value]


def _shaded(ws):
    """The headings marked as required, by their shading."""
    return {str(c.value or "").strip() for row in ws.iter_rows() for c in row
            if c.value and c.fill is not None
            and getattr(c.fill.fgColor, "rgb", "") .endswith(rt.REQUIRED)}


def test_the_counter_form_asks_for_what_the_report_needs(tmp_path):
    """Four counters record no payment channel, 85% of enquiries name nobody,
    and half the tasks are untimed. The form now asks for all three."""
    out = rt.build_counter_template(tmp_path / "form.xlsx", day=date(2026, 9, 1))
    ws = load_workbook(out).active
    marked = _shaded(ws)
    for needed in ("Cash", "bKash / Nagad", "Card", "Cheque",
                   "Bank / Online transfer", "Received by",
                   "Time spent (minutes)", "Sales Amount (BDT)"):
        assert needed in marked, needed
    text = " | ".join(_labels(ws))
    for block in ("Ticket Issue", "Ticket Reissue", "Ticket Refund"):
        assert block in text
    assert "PNR Booked from" in marked


def test_the_visit_form_asks_for_the_three_fields_it_keeps_losing(tmp_path):
    out = rt.build_visit_template(tmp_path / "form.xlsx", day=date(2026, 9, 1))
    ws = load_workbook(out).active
    marked = _shaded(ws)
    for needed in ("Date & Day", "Zone", "Name",
                   "Productivity / Monthly Sale", "Name of Agent"):
        assert needed in marked, needed
    # and it says how to write the figure, because a range is discarded
    assert "3-5 lakh" in " | ".join(_labels(ws))


def test_a_required_field_is_marked_by_shading_and_not_by_a_mark(tmp_path):
    """A trailing ' *' on a heading was enough to send 'Contact No' down a
    looser match and land every phone number in the contact-person column.
    Required is shown by colour, so the headings stay exactly what the report
    already knows how to read."""
    for build in (rt.build_counter_template, rt.build_visit_template):
        ws = load_workbook(build(tmp_path / "f.xlsx")).active
        assert _shaded(ws)
        for label in _labels(ws):
            assert not label.strip().endswith("*"), label


def test_the_form_the_counters_get_is_the_layout_they_already_use(tmp_path):
    """A form nobody recognises is a form nobody fills in."""
    out = rt.build_counter_template(tmp_path / "form.xlsx", day=date(2026, 9, 4))
    ws = load_workbook(out).active
    text = " | ".join(_labels(ws))
    assert "Counter Daily Activities of 4 Sep 2026" in text
    for familiar in ("Counter Name", "Action", "Count PNR", "Employee Name",
                     "Employee ID", "PNR", "Customer Mobile No"):
        assert familiar in text, familiar


@pytest.mark.parametrize("build", [rt.build_counter_template,
                                   rt.build_visit_template])
def test_a_blank_form_is_not_read_as_a_month_of_work(tmp_path, build):
    """The prefilled Action and serial columns must not add up to anything."""
    out = build(tmp_path / "form.xlsx")
    assert load_workbook(out).active.freeze_panes    # header stays while scrolling
    if build is rt.build_counter_template:
        rows, _, _ = cm.parse_workbook(out, month=9)
        assert rows == []
    else:
        assert vm.parse_workbook(out, month=9, year=2026).visits == []


def test_the_day_total_survives_in_the_form(tmp_path):
    """Those typed totals are the only independent check the report has: the
    parse is verified against them 29 times out of 29. Dropping the line to
    save the counter a sum would remove the proof."""
    out = rt.build_counter_template(tmp_path / "form.xlsx")
    text = " | ".join(_labels(load_workbook(out).active))
    for line in ("Previous Sales :", "Total Sales :", "Refund :",
                 "Net Sell After Refund :"):
        assert line in text, line


def test_the_form_asks_for_every_channel_the_parser_can_credit(tmp_path):
    """If the report grows a new required column, the form must grow it too --
    otherwise the gap it exists to close quietly reopens."""
    out = rt.build_counter_template(tmp_path / "form.xlsx")
    offered = " ".join(_labels(load_workbook(out).active)).lower()
    for channel in ("cash", "bkash", "card", "cheque", "bank"):
        assert channel in offered, channel


# --- the round trip: a form we hand out has to be one we can read back ------

def _fill_counter_form(path, out):
    """Type one synthetic sale onto the blank form, the way a counter would."""
    wb = load_workbook(path)
    ws = wb.active
    hdr = next(r[0].row for r in ws.iter_rows()
               if cm._is_sales_header([cm._hkey(c.value) for c in r]))
    cols = {cm._hkey(c.value).replace(" ", ""): c.column
            for c in ws[hdr] if c.value}
    for key, value in (("employeename", "SYNTHETIC PERSON"),
                       ("employeeid", "USBA-9999"), ("pnr", "1ZZ999"),
                       ("salesamount", 12345), ("cash", 12345)):
        col = next((c for k, c in cols.items() if k.startswith(key)), None)
        assert col, key
        ws.cell(row=hdr + 1, column=col, value=value)
    wb.save(out)
    return out


def test_a_sale_typed_on_the_counter_form_is_read_back(tmp_path):
    """The form is only worth handing out if the report can read it.

    The report decides issue/reissue/refund from the Action cell ON THE ROW,
    so a form that left that column blank would have produced a sheet whose
    every sale was silently discarded.
    """
    blank = rt.build_counter_template(tmp_path / "blank.xlsx",
                                      day=date(2026, 9, 5))
    filled = _fill_counter_form(
        blank, tmp_path / "Counter Daily Activities of 05 Sep 2026.xlsx")
    rows, _, meta = cm.parse_workbook(filled, month=9)
    got = [r for r in rows if r.amount == 12345]
    assert got, f"nothing read back from the form we hand out ({len(rows)} rows)"
    row = got[0]
    assert row.block == "ISSUE"          # from the prefilled Action column
    assert row.emp_id == "USBA-9999"
    assert row.emp_name == "SYNTHETIC PERSON"
    assert row.pnr == "1ZZ999"
    assert row.payments.get("Cash") == 12345
    assert row.pay_status == "ok"
    assert row.day == 5                  # from the form's own title line


def test_a_visit_typed_on_the_visit_form_is_read_back(tmp_path):
    """A visit row with no serial in the first column is not read as a visit,
    so the form writes the serials in."""
    blank = rt.build_visit_template(tmp_path / "blank.xlsx",
                                    day=date(2026, 9, 5))
    wb = load_workbook(blank)
    ws = wb.active
    vh = next(r[0].row for r in ws.iter_rows()
              if vm._is_block_header([vm._key(c.value) for c in r]))
    for row in ws.iter_rows(min_row=1, max_row=vh):
        for c in row:
            key = vm._key(c.value)
            if key.startswith("date"):
                ws.cell(row=c.row, column=c.column + 1, value="05/09/2026")
            elif key.startswith("zone"):
                ws.cell(row=c.row, column=c.column + 1, value="Zone 7")
            elif key == "name":
                ws.cell(row=c.row, column=c.column + 1, value="SYNTHETIC REP")
    cols = {vm._column_of(c.value): c.column for c in ws[vh] if c.value}
    for field, value in (("agency", "SYNTHETIC TRAVELS LTD"),
                         ("phone", "01700000000"),
                         ("productivity", "BDT 300000")):
        assert cols.get(field), field
        ws.cell(row=vh + 1, column=cols[field], value=value)
    out = tmp_path / "Visit Report Sep 2026.xlsx"
    wb.save(out)

    data = vm.parse_workbook(out, month=9, year=2026)
    got = [v for v in data.visits if "SYNTHETIC" in v.agency.upper()]
    assert got, f"nothing read back ({len(data.visits)} visits)"
    visit = got[0]
    assert visit.rep == "SYNTHETIC REP"
    assert visit.day == date(2026, 9, 5)
    assert visit.zone == "Zone 7"
    assert visit.phone == "01700000000"       # not the contact-person column
    assert visit.productivity == 300000


def test_the_visit_forms_headings_route_to_the_fields_they_name(tmp_path):
    """'Contact No' means the phone number. It stopped meaning that the moment
    a ' *' was appended to it."""
    ws = load_workbook(rt.build_visit_template(tmp_path / "f.xlsx")).active
    vh = next(r[0].row for r in ws.iter_rows()
              if vm._is_block_header([vm._key(c.value) for c in r]))
    routed = {vm._column_of(c.value) for c in ws[vh] if c.value}
    for field in ("agency", "contact", "phone", "productivity", "routes",
                  "remarks", "location", "designation"):
        assert field in routed, field
