"""The decline list joined to the contacts in the visit reports.

Fixtures are synthetic: the real visit reports carry agency contacts, staff
names and mobile numbers.
"""
from datetime import date

from openpyxl import load_workbook

from src import sales_call_list as scl
from src import sales_movement as sm

AUG = sm.Settings(period_from=date(2026, 8, 1), period_to=date(2026, 8, 31),
                  trailing=3, threshold=0.20)


class _Visit:
    def __init__(self, agency, *, contact="", designation="", phone="",
                 rep="", day=None):
        self.agency, self.contact = agency, contact
        self.designation, self.phone = designation, phone
        self.rep, self.day = rep, day


class _Data:
    def __init__(self, visits):
        self.visits = visits


def _row(customer, window, amount, **meta):
    base = {"customer": customer, "window": window, "amount": amount,
            "tickets": 5, "customer_id": f"C-{customer}", "iata": "1234567",
            "zone": "Zone-4", "station": "DAC", "sales_person": "Rep A",
            "agent_type": "BSP", "channel": "AGENCY"}
    base.update(meta)
    return base


def _result():
    rows = []
    rows += [_row("AB Travels", 0, 1_000_000,
                  last_bought=date(2026, 8, 28))] + \
        [_row("AB Travels", i, 9_000_000) for i in (1, 2, 3)]
    rows += [_row("Quiet Tours", i, 3_000_000,
                  last_bought=date(2026, 5, 2)) for i in (1, 2, 3)]
    rows += [_row("Small Dip", 0, 100_000,
                  last_bought=date(2026, 8, 20))] + \
        [_row("Small Dip", i, 900_000) for i in (1, 2, 3)]
    rows += [_row("Growing", 0, 9_000_000)] + \
        [_row("Growing", i, 1_000_000) for i in (1, 2, 3)]
    return sm.build(rows, AUG)


CONTACTS = _Data([
    _Visit("AB Travel", contact="Mr Rahman", designation="Manager",
           phone="01711000000", rep="Karim", day=date(2026, 7, 10)),
    _Visit("Quiet Tours Ltd.", contact="Ms Akter", phone="01811000000",
           rep="Nadia", day=date(2026, 6, 4)),
])


# --------------------------------------------------------------------------
# matching
# --------------------------------------------------------------------------
def test_an_agency_is_matched_across_how_the_name_was_typed():
    """The sheet says 'AB Travels', the rep wrote 'AB Travel'."""
    got = scl.contacts_from_visits(CONTACTS)
    rows = scl.build(_result(), got)
    ab = next(r for r in rows if r.movement.customer == "AB Travels")
    assert ab.contact.phone == "01711000000"
    assert ab.contact.contact == "Mr Rahman"


def test_a_trade_suffix_does_not_stop_a_match():
    rows = scl.build(_result(), scl.contacts_from_visits(CONTACTS))
    q = next(r for r in rows if r.movement.customer == "Quiet Tours")
    assert q.contact.phone == "01811000000"


def test_a_later_visit_wins_but_never_erases_a_number_with_a_blank():
    data = _Data([
        _Visit("AB Travel", contact="Mr Rahman", phone="01711000000",
               day=date(2026, 6, 1)),
        _Visit("AB Travel", contact="Mr Khan", phone="", day=date(2026, 8, 1)),
    ])
    got = scl.contacts_from_visits(data)["abtravel"]
    assert got.contact == "Mr Khan"           # newer name taken
    assert got.phone == "01711000000"         # older number kept
    assert got.last_seen == date(2026, 8, 1)


# --------------------------------------------------------------------------
# who is on the list
# --------------------------------------------------------------------------
def test_only_the_agencies_that_went_backwards_are_listed():
    names = [r.movement.customer
             for r in scl.build(_result(), scl.contacts_from_visits(CONTACTS))]
    assert "Growing" not in names
    assert set(names) == {"AB Travels", "Quiet Tours", "Small Dip"}


def test_the_list_is_ranked_by_money_lost():
    rows = scl.build(_result(), scl.contacts_from_visits(CONTACTS))
    assert [r.movement.customer for r in rows][0] == "AB Travels"
    assert rows[0].movement.change < rows[-1].movement.change


def test_an_agency_with_no_contact_is_kept_not_dropped():
    """An account losing real money that nobody has visited is the most
    important row on the sheet, not one to hide."""
    rows = scl.build(_result(), scl.contacts_from_visits(CONTACTS))
    dip = next(r for r in rows if r.movement.customer == "Small Dip")
    assert dip.reachable is False
    assert dip.contact.phone == ""


def test_it_works_with_no_visit_reports_at_all():
    rows = scl.build(_result(), None)
    assert len(rows) == 3
    assert all(r.reachable is False for r in rows)


# --------------------------------------------------------------------------
# what each row says
# --------------------------------------------------------------------------
def test_days_silent_counts_to_the_end_of_the_period():
    rows = scl.build(_result(), {})
    q = next(r for r in rows if r.movement.customer == "Quiet Tours")
    assert q.days_silent == (date(2026, 8, 31) - date(2026, 5, 2)).days


def test_the_reason_names_the_bucket_in_plain_words():
    rows = {r.movement.customer: r
            for r in scl.build(_result(), {})}
    assert rows["Quiet Tours"].why == "Stopped buying altogether"
    assert "Down 89%" in rows["AB Travels"].why


def test_a_refunder_is_described_as_one():
    rows = [_row("Refunder", 0, -200_000)] + \
        [_row("Refunder", i, 2_000_000) for i in (1, 2, 3)]
    got = scl.build(sm.build(rows, AUG), {})
    assert got[0].why == "Refunded more than they bought"


# --------------------------------------------------------------------------
# the workbook
# --------------------------------------------------------------------------
def _book(tmp_path, contacts=None):
    out = tmp_path / "calls.xlsx"
    scl.build_workbook(_result(), out,
                       contacts if contacts is not None
                       else scl.contacts_from_visits(CONTACTS))
    return load_workbook(out)


def _text(ws) -> str:
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row
                    if c.value is not None)


def _column(ws, header, limit=40):
    for r in range(1, ws.max_row + 1):
        for c in range(1, 22):
            if ws.cell(row=r, column=c).value == header:
                return [ws.cell(row=i, column=c).value
                        for i in range(r + 1, min(r + 1 + limit,
                                                  ws.max_row + 1))]
    raise AssertionError(f"no column headed {header!r}")


def test_the_workbook_has_one_sheet_a_rep_can_work_from(tmp_path):
    wb = _book(tmp_path)
    assert wb.sheetnames == ["Call list"]
    assert _column(wb["Call list"], "Phone")[0] == "01711000000"


def test_the_workbook_leaves_blank_columns_for_the_rep_to_fill(tmp_path):
    ws = _book(tmp_path)["Call list"]
    for header in ("Called on", "Spoke to", "Outcome"):
        assert _column(ws, header)[0] is None


def test_the_header_says_how_many_can_actually_be_rung(tmp_path):
    said = _text(_book(tmp_path)["Call list"])
    assert "2 of 3 have a phone number" in said.replace("\n", " ")


def test_a_contact_name_without_a_number_is_not_someone_to_ring():
    """'phone or contact' counted a name on its own as reachable -- a caption
    this data happened to satisfy today and would not have tomorrow."""
    got = scl.Contact(agency="A", contact="Mr Rahman", phone="")
    assert got.reachable is False
    assert got.named is True
    assert scl.Contact(agency="A", phone="01711000000").reachable is True


def test_an_empty_list_says_so_rather_than_printing_a_bare_grid(tmp_path):
    rows = [_row("Growing", 0, 9_000_000)] + \
        [_row("Growing", i, 1_000_000) for i in (1, 2, 3)]
    out = tmp_path / "empty.xlsx"
    scl.build_workbook(sm.build(rows, AUG), out, {})
    ws = load_workbook(out)["Call list"]
    assert "Nobody went backwards this period." in _text(ws)


def test_the_filename_names_the_period():
    assert scl.default_filename(AUG) == "Call_List_Aug2026.xlsx"


# --------------------------------------------------------------------------
# a name match is not an identity -- 80 keys in this warehouse cover two or
# more accounts, so a silent match can send a rep to the wrong company
# --------------------------------------------------------------------------
def _twins():
    rows = []
    for cid in ("C-1", "C-2"):
        rows += [_row("Sky Travels", 0, 1_000_000, customer_id=cid)] + \
            [_row("Sky Travels", i, 9_000_000, customer_id=cid)
             for i in (1, 2, 3)]
    rows += [_row("Solo Tours", 0, 100_000, customer_id="C-3")] + \
        [_row("Solo Tours", i, 900_000, customer_id="C-3") for i in (1, 2, 3)]
    return sm.build(rows, AUG)


def test_two_accounts_under_one_name_are_both_marked_for_checking():
    rows = scl.build(_twins(), {})
    sky = [r for r in rows if r.movement.customer == "Sky Travels"]
    assert len(sky) == 2
    assert all(r.ambiguous for r in sky)


def test_an_unambiguous_agency_is_not_marked():
    rows = scl.build(_twins(), {})
    solo = next(r for r in rows if r.movement.customer == "Solo Tours")
    assert solo.ambiguous is False


def test_the_reason_column_tells_the_rep_to_check(tmp_path):
    out = tmp_path / "twins.xlsx"
    scl.build_workbook(_twins(), out, {})
    ws = load_workbook(out)["Call list"]
    why = [v for v in _column(ws, "Why they are on this list") if v]
    assert any("CHECK THE CONTACT" in str(v) for v in why)


def test_the_sheet_shows_which_agency_the_contact_was_filed_under(tmp_path):
    """'AB Travel' in the visit report against 'AB Travels' in the sales
    system is a good match; seeing the other name is how a rep spots a bad
    one."""
    ws = _book(tmp_path)["Call list"]
    filed = [v for v in _column(ws, "Contact is filed under") if v]
    assert "AB Travel" in filed


def test_the_ambiguous_count_is_on_the_header_tiles(tmp_path):
    out = tmp_path / "twins2.xlsx"
    scl.build_workbook(_twins(), out, {})
    said = _text(load_workbook(out)["Call list"])
    assert "CHECK THE CONTACT" in said.upper()


def test_the_call_list_covers_fallers_whatever_the_direction_filter_says():
    """The grid can be showing risers; the call list is inherently about
    the agencies going backwards, so it must not quietly follow `direction`
    and come out empty."""
    rows = []
    rows += [_row("Falling", 0, 1_000_000)] + \
        [_row("Falling", i, 9_000_000) for i in (1, 2, 3)]
    rows += [_row("Rising", 0, 9_000_000)] + \
        [_row("Rising", i, 1_000_000) for i in (1, 2, 3)]
    up_only = sm.Settings(period_from=AUG.period_from, period_to=AUG.period_to,
                          trailing=3, threshold=0.20, direction=sm.INCREASED)
    res = sm.build(rows, up_only)
    assert [m.customer for m in res.reported] == ["Rising"]      # the grid
    assert [r.movement.customer for r in scl.build(res, {})] == ["Falling"]


# --------------------------------------------------------------------------
# a grouped agency opens to its accounts here too
# --------------------------------------------------------------------------
class _M:
    """A movement carrying accounts, as sales_movement now produces."""

    def __init__(self):
        self.customer, self.customer_id = "BE FRESH LIMITED", "10000277"
        self.iata = self.zone = self.station = self.sales_person = ""
        self.bucket, self.baseline, self.current = sm.DECLINED, 900.0, 100.0
        self.change_pct, self.last_bought = -0.89, date(2026, 8, 20)
        self.accounts = [("10000277", "BE FRESH LIMITED", 60.0, 500.0),
                         ("11662412", "Be Fresh Limited (IATA)", 40.0, 400.0)]
        self.is_group = True

    @property
    def change(self):
        return self.current - self.baseline


class _Res:
    settings = AUG
    declined = [_M()]
    lapsed: list = []
    refunded: list = []


def test_the_accounts_open_under_the_agency_on_the_call_list(tmp_path):
    out = tmp_path / "calls.xlsx"
    scl.build_workbook(_Res(), out, {})
    ws = load_workbook(out)["Call list"]
    header = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "Agency":
            header = r
    assert str(ws.cell(row=header + 1, column=1).value).startswith(
        "BE FRESH LIMITED")
    child = header + 2
    assert str(ws.cell(row=child, column=1).value).strip() == "BE FRESH LIMITED"
    assert ws.cell(row=child, column=2).value == "10000277"
    assert ws.row_dimensions[child].outlineLevel == 1
    assert ws.row_dimensions[child].hidden is True
    assert ws.sheet_properties.outlinePr.summaryBelow is False
