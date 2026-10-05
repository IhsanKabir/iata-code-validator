"""The read-only season check, against a stand-in Zenith.

Every page is built in the layout captured from the live screens (Fare
Editor, Quick Sale, penalty and tax pages), with invented fares.
"""
import json
import threading
from datetime import date

import pytest

from src import pricing_check as pc
from src import pricing_client as cl
from src import pricing_parse as pp

TODAY = date(2026, 10, 5)
SEASONS = [{"id": 34198, "group": "Summer 2026", "start": "2026/03/28",
            "end": "2026/10/24"},
           {"id": 34199, "group": "Winter 2026-27", "start": "2026/10/24",
            "end": "2027/03/27"},
           {"id": 42946, "group": "Summer 2027", "start": "2027/03/27",
            "end": "2027/10/30"}]
LEGS = {"68": {"Text": "DAC-CCU"}, "69": {"Text": "CCU-DAC"},
        "244": {"Text": "DAC-DXB"}, "245": {"Text": "DXB-DAC"}}
AIRPORTS = [{"Id": 6707, "Text": "DAC - Dhaka"},
            {"Id": 6261, "Text": "CCU - Kolkata"},
            {"Id": 6942, "Text": "DXB - Dubai Intl"}]
PAX = {"1": {"Code": "AD", "Text": "Adult(s)"},
       "2": {"Code": "CH", "Text": "Child(ren)"}}


# ---- the safety rule ---------------------------------------------------------

BASE = "https://z.test"


@pytest.mark.parametrize("method,path,data", [
    ("POST", cl.FARE_EDITOR + "/SaveFares", None),
    ("POST", "/TTIDotNet/Transport/TransportNetBO2/Pricing/"
             "FareEditorImportExcel.aspx", None),
    ("POST", cl.PENALTIES, {"Action": "Appliquer", "Made": "Update"}),
    ("POST", cl.PENALTIES, {"Action": "Appliquer", "Made": "Seek",
                            "penMod": "300"}),
    ("POST", cl.PENALTIES, {"Action": "Delete", "Made": "Seek"}),
    ("POST", cl.TAXES, {"hAction": "Update"}),
    ("POST", cl.TAXES, {"hAction": "Delete"}),
    ("GET", "/newui/aerien/commercial/anything.asp", None),
])
def test_anything_that_could_change_zenith_is_refused(method, path, data):
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed(method, BASE + path, data)


@pytest.mark.parametrize("command", [
    "BS1Y2", "TKT", "VOID", "XP", "XS1", "FOP CASH", "RC", "PRS",
    "FDIDACCCU/D12FEB26/R,NP;TKT", "FDIDACCCU", "HE", "",
    "FDIDACCCU/D12FEB/R,NP",              # no year: Quick Sale rolls it on
])
def test_only_dated_fare_displays_reach_quick_sale(command):
    from urllib.parse import quote
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed("GET", f"{BASE}{cl.QS_COMMAND}?taskId=x&txtRequest="
                         f"{quote(command)}")


@pytest.mark.parametrize("command", [
    "FDIDACCCU/D20OCT26/R,NP,WT", "FDIDACCCU/D20OCT26/R,NP",
    "FDIDXBDAC/D27FEB27/R,NP,WT,FC-AED", "FDIDACDXB/D12NOV26"])
def test_the_fare_displays_the_check_uses_are_allowed(command):
    from urllib.parse import quote
    cl.check_allowed("GET", f"{BASE}{cl.QS_COMMAND}?taskId=x&txtRequest="
                     f"{quote(command)}")


def test_the_reads_the_check_uses_are_allowed():
    cl.check_allowed("POST", BASE + cl.LOAD_FARES)
    cl.check_allowed("POST", BASE + cl.FARE_BASIS_LIST)
    cl.check_allowed("GET", BASE + cl.FARE_EDITOR + "?idFarebasis=1")
    cl.check_allowed("POST", BASE + cl.PENALTIES,
                     pp.penalty_search_form([1], [2], [3], [1]))
    cl.check_allowed("POST", BASE + cl.TAXES,
                     pp.tax_search_form(1, 2, "05/10/2026"))
    cl.check_allowed("POST", BASE + cl.TAXES, {"hAction": "Affiche_Result"})


# ---- the stand-in Zenith -----------------------------------------------------

def fe_page():
    return ("<script>var legs = " + json.dumps(LEGS) + ";\nvar periods = "
            + json.dumps(SEASONS) + ";\nvar airports = "
            + json.dumps(AIRPORTS) + ";\nvar passengerTypes = "
            + json.dumps(PAX) + ";</script>")


def qs_page(currency, taxes, season, rows):
    lines = [f"Currency:{currency}&nbsp;/BS&nbsp;DACCCU&nbsp;"
             + ("Taxes included" if taxes else "Without taxes"),
             "| FARE BASIS | OW | RT | PRBD | LVL | REF | MOD | MIN | MAX",
             f"{season}: 28MAR26**24OCT26"]
    for i, (code, ow, rt) in enumerate(rows, 1):
        lines.append(f"{i} | {code} | {ow if ow is not None else '-'} | "
                     f"{rt if rt is not None else '-'} | Y | 8 | PEN | PEN "
                     f"| 0 | -1")
    return "<br/>".join(f"<span>{l}</span>" for l in lines)


def penalty_page():
    def box(kind, i, label, attr):
        return (f'<input type="checkbox" data-{attr}="{i}" '
                f'name="{i}_{kind}" id="{i}_{kind}"><label for="{i}_{kind}">'
                f'<span>{label}</span></label>')
    parts = [box("IDClasse", 11, "YBDCCUO", "idClass"),
             box("IDClasse", 12, "YBDCCUR", "idClass"),
             box("IDLeg", 68, "DAC - CCU", "idLeg"),
             box("IDLeg", 69, "CCU - DAC", "idLeg")]
    parts += [box("IDPeriode", s["id"], s["group"], "idPeriod")
              for s in SEASONS]
    parts += [box("IDTypePassager", 1, "AD - Adult(s)", "idPassengerType")]
    return "".join(parts)


def penalty_results(rows):
    def td(col, label, value):
        return (f'<td class="col-{col}"><div class="visible-xs-inline '
                f'xs-left">{label}</div><div class="visible-xs-inline '
                f'xs-right">{value}</div></td>')
    body = [
        "<tr>" + td("classe", "Class", r[0]) + td("leg", "Leg", r[1])
        + td("typpass", "Pass. Type", r[2]) + td("periode", "Seasons", r[3])
        + td("interval", "Validity", r[4])
        + td("penalitemodif", "Modification penalty estimate", r[5])
        + td("penaliteremb", "", "")
        + td("pendevremb", "Refund penalty estimate", r[6]) + "</tr>"
        for r in rows]
    # Zenith shows a big answer as several tables of 100 rows
    tables = "".join(
        '<table class="previewMod"><thead><tr><th class="col-classe">Class'
        '</th><th>Modif. penalty</th></tr></thead>' + "".join(body[i:i + 100])
        + "</table>" for i in range(0, max(len(body), 1), 100))
    return f"<html><body>{tables}</body></html>"


def tax_page():
    return ('<select name="RecCodeDestGeo"><option value="">-- Airports --'
            '</option><option value="6707">DAC - Dhaka</option><option '
            'value="6261">CCU - Kolkata</option><option value="6942">DXB - '
            'Dubai Intl</option></select>')


def tax_list(rows):
    head = "".join(f"<th>{h}</th>" for h in pp._TAX_HEADER)
    body = ""
    for name, code, dep, arr, value, intl in rows:
        body += ("<tr>" + "".join(
            f"<td>{lab} {val}</td>" for lab, val in zip(
                pp._TAX_HEADER[:5], (name, code, dep, arr, value)))
            + f"<td>International {'&#xe5ca;' if intl else ''}</td>"
            + "<td>Domestic </td><td>History History</td></tr>")
    return f"<table><tr>{head}</tr>{body}</table>"


def fare(leg, pax, cur, ow, rt=0, bag=20, excess=5):
    return {"IdLeg": leg, "IdPassengerType": pax, "Currency": cur,
            "AmountOW": ow, "AmountRT": rt, "AmountYQ": 0.0, "AmountYR": 0.0,
            "BaggageWeightAllowance": bag, "BaggageExcessWeightAmount": excess,
            "BagPieceAllowance": 0}


class FakeZenith:
    """Fares per (fare basis id, season id); the rest per season."""

    def __init__(self):
        self.fares = {}
        self.qs = {}            # (od, season name, taxes, currency) -> rows
        self.taxes = {}         # season day -> rows
        self.penalties = []     # (class, leg, pax, season, window, mod, ref)
        self.sent = []

    def fare_basis_list(self):
        return [{"IdFareBasis": 11, "Code": "YBDCCUO", "Text": "Y one way",
                 "IsEnabled": True, "BookingClassCode": "Y"},
                {"IdFareBasis": 12, "Code": "YBDCCUR", "Text": "Y return",
                 "IsEnabled": True, "BookingClassCode": "Y"},
                {"IdFareBasis": 13, "Code": "OLDX", "Text": "off",
                 "IsEnabled": False, "BookingClassCode": "X"}]

    def fare_editor_page(self, _id):
        return fe_page()

    def load_fares(self, fb, season):
        return [dict(r) for r in self.fares.get((fb, season), [])]

    def quick_sale(self, command):
        cl.check_allowed("GET", f"{BASE}{cl.QS_COMMAND}?txtRequest="
                         + command.replace(",", "%2C").replace("/", "%2F"))
        self.sent.append(command)
        od = command[3:9]
        day = command.split("/D")[1][:7]
        season = {"13OCT26": "SUMMER 2026",
                  "31OCT26": "WINTER 2026-27"}[day]
        taxes = ",WT" not in command
        cur = command.split("FC-")[1] if "FC-" in command else "BDT"
        rows = self.qs.get((od, season, taxes, cur), [])
        return qs_page(cur, taxes, season, rows)

    def tax_page(self):
        return tax_page()

    def tax_search(self, form):
        cl.check_allowed("POST", BASE + cl.TAXES, form)
        return tax_list(self.taxes.get(form["RecEffDate"], []))

    def penalty_page(self):
        return penalty_page()

    def penalty_search(self, form):
        cl.check_allowed("POST", BASE + cl.PENALTIES, form)
        self.sent.append(form)
        seasons = {s["group"] for s in SEASONS
                   if f"{s['id']}_IDPeriode" in form}
        return penalty_results([p for p in self.penalties
                                if p[3] in seasons])


def same_both_seasons(z):
    for season in (34198, 34199):
        z.fares[(11, season)] = [fare(68, 1, "USD", 300), fare(68, 2, "USD",
                                                               225)]
        z.fares[(12, season)] = [fare(68, 1, "USD", 0, rt=500)]
    for season in ("SUMMER 2026", "WINTER 2026-27"):
        z.qs[("DACCCU", season, False, "BDT")] = [("YBDCCUO", 37097, None),
                                                  ("YBDCCUR", None, 61828)]
        z.qs[("DACCCU", season, True, "BDT")] = [("YBDCCUO", 43100, None),
                                                 ("YBDCCUR", None, 72443)]
        z.qs[("DACCCU", season, False, "USD")] = [("YBDCCUO", 300, None),
                                                  ("YBDCCUR", None, 500)]
    for day in ("13/10/2026", "31/10/2026"):
        z.taxes[day] = [("Embarkation Fee", "BD", "Dhaka", "Kolkata",
                         "500 BDT", True),
                        ("Q - Carrier Imposed Fee", "YQ", "Dhaka", "Kolkata",
                         "5 USD", True)]
    for season in ("Summer 2026", "Winter 2026-27"):
        z.penalties += [("Y one way", "DAC - CCU", "Adult(s)", season,
                         "0 - 72 hour(s) before flight departure", "300 USD",
                         "480 USD")]
    return z


def index_for(tmp_path):
    idx = pc.RouteIndex(tmp_path / "routes.json")
    idx.legs = {11: {"DAC-CCU"}, 12: {"DAC-CCU"}}
    return idx


def run(z, tmp_path, **kw):
    settings = pc.Settings(routes_text="DAC-CCU", include_return=False, **kw)
    return pc.run_check(z, index_for(tmp_path), settings, today=TODAY)


# ---- reading the screens -----------------------------------------------------

def test_the_fare_editor_lookups_are_read():
    got = pp.parse_fare_editor_page(fe_page())
    assert got.legs[68] == "DAC-CCU"
    assert got.airport_code["DXB"] == 6942
    assert got.passenger_types[1] == "AD - Adult(s)"
    assert [s.name for s in got.seasons][:2] == ["Summer 2026",
                                                 "Winter 2026-27"]


def test_a_quick_sale_display_is_read_and_a_changed_one_refused():
    shown = pp.parse_quick_sale(qs_page("USD", True, "SUMMER 2026",
                                        [("YBDCCUO", "1,421.87", None)]))
    assert (shown.currency, shown.taxes_included) == ("USD", True)
    assert shown.fares[0].ow == 1421.87 and shown.fares[0].rt is None
    with pytest.raises(pp.LayoutChanged):
        pp.parse_quick_sale(qs_page("USD", True, "SUMMER 2026", [])
                            .replace("| MAX", "| MAX | NEW"))
    err = pp.parse_quick_sale("Following command was not recognized : 'I'.")
    assert err.error and not err.fares


def test_penalty_and_tax_pages_are_read():
    opts = pp.parse_penalty_page(penalty_page())
    assert opts.classes[11] == "YBDCCUO" and opts.legs[68] == "DAC - CCU"
    rows = pp.parse_penalty_results(penalty_results(
        [("Y", "DAC - CCU", "Adult(s)", "Summer 2026", "0 - 72 h", "300 USD",
          "480 USD")]))
    assert rows[0].change == "300 USD" and rows[0].refund == "480 USD"
    assert pp.parse_tax_airports(tax_page()) == {"DAC": 6707, "CCU": 6261,
                                                 "DXB": 6942}
    taxes = pp.parse_tax_list(tax_list([("BD fee", "BD", "Dhaka", "Kolkata",
                                         "500 BDT", True)]))
    assert taxes[0].code == "BD" and taxes[0].value == "500 BDT"
    assert taxes[0].international and not taxes[0].domestic


# ---- seasons, dates, routes --------------------------------------------------

def test_current_versus_next_and_all():
    seasons = pp.parse_fare_editor_page(fe_page()).seasons
    base, nxt = pc.seasons_to_compare(seasons, TODAY, "next")
    assert base.name == "Summer 2026" and [s.name for s in nxt] == [
        "Winter 2026-27"]
    _, all_ = pc.seasons_to_compare(seasons, TODAY, "all")
    assert [s.name for s in all_] == ["Winter 2026-27", "Summer 2027"]
    # on the shared boundary day the new season has begun
    base, _ = pc.seasons_to_compare(seasons, date(2026, 10, 24), "next")
    assert base.name == "Winter 2026-27"


def test_a_check_date_is_ahead_inside_the_season_and_carries_the_year():
    s = pp.parse_fare_editor_page(fe_page()).seasons[0]       # ends 24 Oct
    day = pc.check_date(s, TODAY)
    assert TODAY < day < date(2026, 10, 24)
    assert pc.fdi_date(day) == "13OCT26"
    late = pc.check_date(s, date(2026, 10, 22))
    assert late == date(2026, 10, 23)                      # never its last day


def test_routes_typed_by_the_user_are_read_with_their_return():
    legs, unknown = pc.parse_routes("dac-ccu, DAC DXB; XXX-YYY",
                                    ["DAC-CCU", "CCU-DAC", "DAC-DXB",
                                     "DXB-DAC"])
    assert legs == {"DAC-CCU", "CCU-DAC", "DAC-DXB", "DXB-DAC"}
    assert "XXX-YYY" in unknown


# ---- the comparisons ---------------------------------------------------------

def test_identical_seasons_report_nothing(tmp_path):
    res = run(same_both_seasons(FakeZenith()), tmp_path)
    assert res.findings == [] and not res.problems
    assert res.compared["Fares"] == 3 and res.compared["Taxes"] == 2


def test_a_changed_amount_currency_and_baggage_are_each_reported(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.fares[(11, 34199)] = [fare(68, 1, "BDT", 310, bag=15),
                            fare(68, 2, "USD", 225)]
    res = run(z, tmp_path)
    changed = {(f.detail, f.value_a, f.value_b) for f in res.findings
               if f.part == "Fares"}
    assert ("Currency", "USD", "BDT") in changed
    assert ("AmountOW", "300", "310") in changed
    assert ("BaggageWeightAllowance", "20", "15") in changed


def test_a_fare_missing_next_season_is_reported(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.fares[(11, 34199)] = [fare(68, 1, "USD", 300)]        # child gone
    res = run(z, tmp_path)
    missing = [f for f in res.findings if f.kind == pc.MISSING]
    assert len(missing) == 1 and "Child" in missing[0].item


def test_a_fare_set_up_but_not_selling_is_caught_by_quick_sale(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.qs[("DACCCU", "WINTER 2026-27", False, "USD")] = [("YBDCCUO", 300,
                                                         None)]
    res = run(z, tmp_path)
    sales = [f for f in res.findings if f.part == "Sales check"]
    assert [(f.kind, f.item) for f in sales] == [("set up, not on sale",
                                                  "YBDCCUR")]


def test_a_tax_changed_between_seasons_is_reported(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.taxes["31/10/2026"][1] = ("Q - Carrier Imposed Fee", "YQ", "Dhaka",
                                "Kolkata", "7 USD", True)
    res = run(z, tmp_path)
    taxes = [f for f in res.findings if f.part == "Taxes"]
    assert [(f.detail, f.value_a, f.value_b) for f in taxes] == [
        ("Value", "5 USD", "7 USD")]


def test_quick_sale_total_tax_is_compared_too(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.qs[("DACCCU", "WINTER 2026-27", True, "BDT")] = [("YBDCCUO", 44000,
                                                        None),
                                                       ("YBDCCUR", None,
                                                        72443)]
    res = run(z, tmp_path)
    qs = [f for f in res.findings if f.part == "QS taxes"]
    assert [(f.item, f.detail) for f in qs] == [("YBDCCUO", "OW tax")]


def test_a_penalty_missing_next_season_is_reported(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.penalties = [p for p in z.penalties if p[3] != "Winter 2026-27"]
    res = run(z, tmp_path)
    pen = [f for f in res.findings if f.part == "Penalties"]
    assert [f.kind for f in pen] == [pc.MISSING]


def test_every_results_table_of_a_big_penalty_answer_is_read():
    rows = [("Y", "DAC - CCU", "Adult(s)", "Summer 2026", f"{i} h", "1 USD",
             "2 USD") for i in range(144)]
    html = penalty_results(rows)
    assert html.count("Modif. penalty") == 2
    assert len(pp.parse_penalty_results(html)) == 144


def test_a_penalty_page_that_stopped_loading_is_not_trusted():
    html = penalty_results([("Y", "DAC - CCU", "Adult(s)", "Summer 2026",
                             "0 h", "1 USD", "2 USD")])
    with pytest.raises(pp.Incomplete):
        pp.parse_penalty_results(html[:-30])


def test_a_penalty_search_that_does_not_load_is_split_then_reported():
    z = same_both_seasons(FakeZenith())
    real = z.penalty_search

    def flaky(form):
        page = real(form)
        classes = [k for k in form if k.endswith("_IDClasse")]
        # two classes at once, or class 12 alone, never finish loading
        return page[:-20] if len(classes) > 1 or "12_IDClasse" in form             else page
    z.penalty_search = flaky
    result = pc.CheckResult()
    rows = pc._penalty_rows(z, [11, 12], [34198, 34199], 68, [1], result,
                            "DAC-CCU")
    searches = [s for s in z.sent if isinstance(s, dict)]
    assert len(searches) == 3                 # the pair, then each alone
    assert len(rows) == 2                     # class 11's rows kept
    assert result.problems == ["DAC-CCU: penalties for fare class 12 did "
                               "not load completely; not checked."]


def test_a_tax_row_of_the_wrong_shape_stops_the_read():
    html = tax_list([("BD fee", "BD", "Dhaka", "Kolkata", "500 BDT", True)])
    broken = html.replace("</table>", "<tr><td>a</td><td>b</td></tr></table>")
    with pytest.raises(pp.LayoutChanged):
        pp.parse_tax_list(broken)
    # a one-cell message line is fine
    ok = html.replace("</table>", "<tr><td>13 taxes</td></tr></table>")
    assert len(pp.parse_tax_list(ok)) == 1


def test_a_fare_quick_sale_shows_as_dash_is_reported(tmp_path):
    z = same_both_seasons(FakeZenith())
    z.qs[("DACCCU", "WINTER 2026-27", False, "USD")] = [
        ("YBDCCUO", 300, None), ("YBDCCUR", None, None)]
    res = run(z, tmp_path)
    sales = [(f.kind, f.item, f.detail, f.value_b) for f in res.findings
             if f.part == "Sales check"]
    assert sales == [("amount differs", "YBDCCUR", "RT USD",
                      "Quick Sale —")]


# ---- redirects never skip the safety rule ------------------------------------

class _Resp:
    def __init__(self, url, status=200, location="", text=""):
        self.url, self.status_code, self.text = url, status, text
        self.headers = {"Location": location} if location else {}


class _Session:
    """A requests session stand-in: answers from a table, records requests."""

    def __init__(self, answers):
        self.answers, self.sent = answers, []

    def get(self, url, **kw):
        assert kw.get("allow_redirects") is False
        self.sent.append(url)
        return self.answers.get(url.split("?")[0].lower(), _Resp(url))

    post = get


def test_a_redirect_to_a_saving_page_is_refused_before_it_is_followed():
    page = (BASE + cl.TAXES).lower()
    sess = _Session({page: _Resp(BASE + cl.TAXES, 302,
                                 cl.FARE_EDITOR + "/SaveFares")})
    client = cl.PricingClient(type("Z", (), {"session": sess})(),
                              base_url=BASE, attempts=1)
    with pytest.raises(cl.ReadOnlyViolation):
        client.tax_page()
    assert all("SaveFares" not in u for u in sess.sent)


def test_quick_sale_opens_through_its_own_redirect():
    tid = "0f8fad5b-d9cb-469f-a165-70867728950e"
    init = (BASE + cl.QS_INIT).lower()
    sess = _Session({init: _Resp(BASE + cl.QS_INIT, 302,
                                 f"{cl.QS_PAGE}?taskId={tid}")})
    client = cl.PricingClient(type("Z", (), {"session": sess})(),
                              base_url=BASE, attempts=1)
    client.quick_sale("FDIDACCCU/D20OCT26/R,NP")
    assert f"taskId={tid}" in sess.sent[-1]


def test_a_redirect_to_sign_in_means_the_session_is_lost():
    page = (BASE + cl.TAXES).lower()
    sess = _Session({page: _Resp(BASE + cl.TAXES, 302,
                                 "https://z.test/otds/login")})
    client = cl.PricingClient(type("Z", (), {"session": sess})(),
                              base_url=BASE, attempts=1)
    with pytest.raises(cl.SessionLost):
        client.tax_page()


def test_without_a_route_list_the_check_says_to_build_one(tmp_path):
    z = same_both_seasons(FakeZenith())
    empty = pc.RouteIndex(tmp_path / "none.json")
    with pytest.raises(ValueError, match="route list"):
        pc.run_check(z, empty, pc.Settings(routes_text="DAC-CCU"),
                     today=TODAY)


def test_the_route_list_is_built_from_enabled_fare_bases_and_kept(tmp_path):
    z = same_both_seasons(FakeZenith())
    idx = pc.RouteIndex(tmp_path / "routes.json")
    pc.rebuild_index(z, idx, today=TODAY)
    assert idx.legs == {11: {"DAC-CCU"}, 12: {"DAC-CCU"}}
    again = pc.RouteIndex(tmp_path / "routes.json")
    assert again.load() and again.fare_bases_for(["DAC-CCU"]) == [11, 12]


def test_stop_ends_the_check(tmp_path):
    z = same_both_seasons(FakeZenith())
    stop = threading.Event()
    stop.set()
    with pytest.raises(InterruptedError):
        pc.run_check(z, index_for(tmp_path), pc.Settings(
            routes_text="DAC-CCU", include_return=False), today=TODAY,
            stop=stop)


# ---- the workbook ------------------------------------------------------------

def test_the_workbook_lists_every_finding_and_what_was_not_checked(tmp_path):
    from openpyxl import load_workbook
    from src import pricing_report as pr
    z = same_both_seasons(FakeZenith())
    z.fares[(11, 34199)] = [fare(68, 1, "BDT", 300), fare(68, 2, "USD", 225)]
    res = run(z, tmp_path)
    res.problems.append("DXB-DAC: Quick Sale said 'x'.")
    path = tmp_path / "check.xlsx"
    pr.build_workbook(res, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Fares", "Sales check", "QS taxes",
                             "Taxes", "Penalties", "Not checked"]
    fares = [[c.value for c in row] for row in wb["Fares"].iter_rows()]
    assert any(r[2] == "Currency" and r[5] == "USD" and r[7] == "BDT"
               for r in fares)
    assert "No differences" in str(wb["Taxes"]["A4"].value)
    assert "Quick Sale said" in str(wb["Not checked"]["A4"].value)
