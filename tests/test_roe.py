"""Exchange-rate (ROE) updates, against a stand-in ROE screen.

The stand-in is built in the live screen's layout (from "ROE SETUP.har"):
a frmDevise form, rows of DeviseTaux{n} / CodeDevise{n} / SAVE_* fields,
ICER boxes carrying value="checked" but NOT ticked, the language select set
by script, page links aff_page_result(n), and a history page per currency.
Rates are invented; the cut-to-6-decimals rule uses the real 7 Oct sheet.
"""
from datetime import date, datetime
from decimal import Decimal

import pytest

from src import pricing_client as cl
from src import roe

TODAY = date(2026, 10, 8)
BASE = "https://z.test"

PAGES = {1: ["AED", "BDT", "EUR", "KWD"], 2: ["NUC", "SAR", "USD"]}
LINKED = {"USD": "E5 ,T6 ", "SAR": ""}


def page_html(nav, state):
    rows = ""
    for n, code in enumerate(PAGES[nav], start=1):
        rate, modified, by, hid = state[code]
        linked = (f' data-linked-taxes="{LINKED[code]}"'
                  if LINKED.get(code) else "")
        rows += (
            f'<tr><td>{code}</td><td><div class="xs-left">Currency</div>'
            f'{code} name</td><td><div>Rate:</div><input type="text" '
            f'name="DeviseTaux{n}" id="DeviseTaux{n}" value="{rate}"'
            f'{linked} onchange="onROEChanged(this)" /></td>'
            f'<td><div>Last modified</div>{modified} GMT</td><td><div>By'
            f'</div>{by}</td><td><a onclick="DisplayHistory({hid})">History'
            f'</a><div>History</div></td><td>ROE scheduled update 0</td><td>'
            f'ROE Auto update - ICER <input name="RoeAutoUpdateEnabled{n}" '
            f'type="checkbox" id="RoeAutoUpdateEnabled{n}" value="checked" />'
            f'<input type="hidden" name="CodeDevise{n}" value="{code}" />'
            f'<input type="hidden" name="SAVE_DeviseTaux{n}" value="{rate}" />'
            f'<input type="hidden" name="SAVE_Subdivision{n}" value="2" />'
            f'<input type="hidden" name="SAVE_RoeAutoUpdateEnabled{n}" '
            f'value="0" /></td></tr>')
    links = "".join(f'<a onclick="aff_page_result({k})">{k}</a>'
                    for k in PAGES)
    return (
        '<form name="frmDevise" action="SaisieTauxDevises.asp" method="post">'
        '<input type="hidden" name="hAction" value="" />'
        '<input type="hidden" name="UpdateLinkedTaxes" value="0" />'
        f'<input type="hidden" name="NAV" value="{nav}" />'
        '<input type="hidden" name="id_societe" value="2035" />'
        '<select name="Id_Langue" class="selectpicker"><option value="">'
        'Language</option><option value="2">English</option></select>'
        '<input type="text" name="CurrencyToSearch" value="" />'
        '<input type="button" value="Save" id="btSave" />'
        f'<table>{rows}</table>{links}</form>'
        "<script>var theForm = document.frmDevise; "
        "theForm.Id_Langue.value = '2';</script></html>")


def history_html(entries):
    rows = "".join(
        f"<tr><td>Modified the {w} GMT</td><td>By {b}</td>"
        f"<td>Rate: {o}</td><td>&#xE5C8;</td><td>Rate: {n}</td></tr>"
        for w, b, o, n in entries)
    return f"<table><tr><th>Modified the</th></tr>{rows}</table>"


class FakeRoe:
    """The ROE screen: saves change its state, history and Quick Sale."""

    def __init__(self):
        self.state = {"AED": ("0.029723", "07/10/2026 03:34:18", "A", 12),
                      "BDT": ("1", "01/09/2018 10:41:50", "A", 1),
                      "EUR": ("0.007222", "07/10/2026 03:34:18", "A", 15),
                      "KWD": ("0.002498", "07/10/2026 03:34:18", "A", 18),
                      "NUC": ("0.008094", "05/10/2026 18:00:03", "A", 2),
                      "SAR": ("0.029866", "07/10/2026 03:34:18", "A", 17),
                      "USD": ("0.008094", "05/10/2026 18:00:03", "A", 3)}
        self.history = {c: [] for c in self.state}
        self.saves = []                 # (nav, form dict)
        self.fail_on_save = set()       # navs whose save raises
        self.reply_differs = {}         # code -> value the reply shows
        self.qs_lag = set()             # codes Quick Sale still shows old

    def _check(self, method, path, form=None, permit=None):
        cl.check_allowed(method, BASE + path, form, permit)

    def roe_page(self):
        self._check("GET", cl.ROE)
        return page_html(1, self.state)

    def roe_post(self, form, permit=None):
        form = list(form)
        self._check("POST", cl.ROE, form, permit)
        f = dict(form)
        nav = int(f["NAV"])
        if f["hAction"] == "Save":
            if nav in self.fail_on_save:
                raise cl.PricingError("Zenith returned 500 3 times.")
            self.saves.append((nav, f))
            for n, code in enumerate(PAGES[nav], start=1):
                new = f[f"DeviseTaux{n}"]
                old, _m, _b, hid = self.state[code]
                if new != old:
                    self.state[code] = (new, "08/10/2026 09:00:00", "Tester",
                                        hid)
                    self.history[code].insert(0, ("08/10/2026 09:00:00",
                                                  "Tester", old, new))
            shown = dict(self.state)
            for code, v in self.reply_differs.items():
                shown[code] = (v,) + shown[code][1:]
            return page_html(nav, shown)
        return page_html(nav, self.state)

    def roe_history(self, hid):
        self._check("GET", cl.ROE_HISTORY + f"?IDDevise={hid}")
        code = next(c for c, v in self.state.items() if v[3] == hid)
        return history_html(self.history[code])

    def quick_sale(self, command):
        from urllib.parse import quote
        self._check("GET", cl.QS_COMMAND + "?taskId=x&txtRequest="
                    + quote(command))
        code = command.rsplit("FC-", 1)[1]
        rate = self.state[code][0]
        if code in self.qs_lag:
            rate = "0.1"
        return (f"<span>Currency:{code}(1 BDT = {Decimal(rate):.8f} {code})"
                f"&nbsp;/BS DACDXB Taxes included</span>")


def sheet(rates, day=TODAY, notes=None):
    return roe.SheetRates(day=day, rates={k: Decimal(v)
                                          for k, v in rates.items()},
                          notes=notes or {})


def plan_for(z, rates, **kw):
    return roe.make_plan(sheet(rates, **kw), roe.read_all_pages(z),
                         today=TODAY)


# ---- the sheet and the rule --------------------------------------------------

def test_the_rate_sheet_is_read_like_the_travelport_one(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Current Date", None, None, "Previous Date", None])
    ws.append([datetime(2026, 10, 7), None, None, datetime(2026, 10, 6)])
    ws.append(["Currency", "Exchange Rate To BDT", None, "Currency", "x"])
    ws.append(["1 KWD", 400.18799, None, "1 KWD", 400.9939])
    ws.append(["1 USD", "123.54\xa0(06-OCT-26)", None, "1 USD", 123.54])
    ws.append(["1 EUR", "n/a", None, "1 EUR", 1])
    path = tmp_path / "07th Oct ROW.xlsx"
    wb.save(path)
    got = roe.read_rate_sheet(path)
    assert got.day == date(2026, 10, 7)
    assert got.rates == {"KWD": Decimal("400.18799"),
                         "USD": Decimal("123.54")}
    assert got.notes == {"USD": "06-OCT-26"}
    assert got.problems == ["EUR: 'n/a' is not a rate."]


@pytest.mark.parametrize("bdt,zenith", [
    # the 7 Oct sheet against what Zenith held before 8 Oct: cut, not round
    ("400.18799", "0.002498"), ("322.09199", "0.003104"),
    ("138.46524", "0.007222"), ("33.643623", "0.029723"),
    ("18.426845", "0.054268"), ("8.011673", "0.124817"),
    ("1.283302", "0.779239"), ("0.368535", "2.713446"),
    ("123.54", "0.008094"),
])
def test_the_sheet_rate_is_turned_round_and_cut_to_six_decimals(bdt, zenith):
    assert roe.zenith_rate(Decimal(bdt)) == Decimal(zenith)


def test_rates_are_written_as_zenith_writes_them():
    assert roe.rate_text(Decimal("0.002490")) == "0.00249"
    assert roe.rate_text(Decimal("1")) == "1"
    assert roe.rate_text(Decimal("2.713446")) == "2.713446"


# ---- reading the screen ------------------------------------------------------

def test_the_form_is_rebuilt_as_the_browser_posts_it():
    form = dict(roe.browser_form(page_html(1, FakeRoe().state), "frmDevise"))
    assert "RoeAutoUpdateEnabled1" not in form      # value="checked", unticked
    assert form["Id_Langue"] == "2"                  # set by the page's script
    assert "btSave" not in form and form["hAction"] == ""
    ticked = roe.browser_form('<form name="f"><input type="checkbox" '
                              'name="a" checked /><input name="b" disabled '
                              'value="x" /></form>', "f")
    assert ticked == [("a", "on")]


def test_the_screen_is_read_page_by_page():
    pages = roe.read_all_pages(FakeRoe())
    assert sorted(pages) == [1, 2]
    usd = pages[2].rows["USD"]
    assert (usd.index, usd.rate, usd.history_id) == (3, Decimal("0.008094"), 3)
    assert usd.linked_taxes == ("E5", "T6")
    assert usd.modified == "05/10/2026 18:00:03"
    assert pages[1].rows["KWD"].linked_taxes == ()


def test_history_and_quick_sale_rates_are_read():
    top = roe.parse_history(history_html([("08/10/2026 03:53:20", "Kabir",
                                           "0.002498", "0.002501")]))
    assert top[0] == roe.HistoryEntry("08/10/2026 03:53:20", "Kabir",
                                      Decimal("0.002498"), Decimal("0.002501"))
    assert roe.quick_sale_rate("<span>Currency:AED(1 BDT = 0.02969600 AED)"
                               " /BS DXBDAC</span>") == ("AED",
                                                         Decimal("0.029696"))


# ---- the plan ----------------------------------------------------------------

def test_the_plan_lists_changes_and_what_is_left_alone():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84", "AED": "33.643623",
                        "USD": "120.5", "XYZ": "1"},
                    notes={"USD": "06-OCT-26"})
    assert [(c.code, roe.rate_text(c.new)) for c in plan.changes] == [
        ("KWD", "0.002501"), ("USD", "0.008298")]
    assert plan.same == ["AED"]
    skipped = dict(plan.skipped)
    assert "not a currency" in skipped["XYZ"]
    assert "left as it is" in skipped["NUC"] and "BDT" in skipped
    assert plan.linked_taxes() == ["E5", "T6"]
    assert any("USD moves +2.52%" in w for w in plan.warnings)   # big change
    assert any("06-OCT-26" in w for w in plan.warnings)


def test_an_old_sheet_is_flagged():
    plan = plan_for(FakeRoe(), {"KWD": "399.84"}, day=date(2026, 10, 7))
    assert any("not today" in w for w in plan.warnings)


# ---- the guard ---------------------------------------------------------------

def test_a_save_without_the_plans_permit_is_refused():
    z = FakeRoe()
    form = dict(roe.browser_form(page_html(1, z.state), "frmDevise"))
    form.update(hAction="Save", DeviseTaux4="0.002501")
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed("POST", BASE + cl.ROE, form)
    # nor a page change that carries a changed rate
    form.update(hAction="")
    with pytest.raises(cl.ReadOnlyViolation, match="changed rates"):
        cl.check_allowed("POST", BASE + cl.ROE, form)
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed("POST", BASE + cl.ROE, {"hAction": "Delete"})


def test_the_permit_accepts_exactly_the_planned_save():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    page = plan.pages[1]
    form = plan.save_form(page, True)
    assert dict(form)["DeviseTaux4"] == "0.002501"
    assert dict(form)["hAction"] == "Save"
    assert dict(form)["UpdateLinkedTaxes"] == "0"    # KWD has no linked tax
    permit = plan.permits(page, True)
    assert permit(form)
    sneaky = dict(form, DeviseTaux1="0.5")
    assert not permit(sneaky)
    assert not permit(dict(form, UpdateLinkedTaxes="1"))


# ---- saving and checking -----------------------------------------------------

def test_rates_are_saved_page_by_page_and_checked_five_ways():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84", "USD": "120.5"})
    res = roe.apply(z, plan, update_linked=True, today=TODAY)
    assert [nav for nav, _f in z.saves] == [1, 2]
    page2 = z.saves[1][1]
    assert page2["UpdateLinkedTaxes"] == "1"         # USD carries E5, T6
    assert z.saves[0][1]["UpdateLinkedTaxes"] == "0"
    assert z.state["KWD"][0] == "0.002501" and z.state["USD"][0] == "0.008298"
    assert z.state["AED"][0] == "0.029723"           # untouched
    assert res.untouched_ok and not res.failed
    for chk in res.checks:
        assert chk.ok, chk
        assert chk.history.startswith("✓") and chk.quick_sale.startswith("✓")


def test_linked_taxes_can_be_left_alone():
    z = FakeRoe()
    plan = plan_for(z, {"USD": "120.5"})
    roe.apply(z, plan, update_linked=False, today=TODAY)
    assert z.saves[0][1]["UpdateLinkedTaxes"] == "0"


def test_a_rate_that_moved_since_the_preview_stops_the_save():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    z.state["EUR"] = ("0.0071", "08/10/2026 08:00:00", "Someone", 15)
    res = roe.apply(z, plan, today=TODAY)
    assert not z.saves
    assert "changed since the preview (EUR)" in res.failed
    assert res.checks[0].saved == "✗ not saved"


def test_a_failed_page_is_reported_and_the_saved_one_kept():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84", "USD": "120.5"})
    z.fail_on_save = {2}
    res = roe.apply(z, plan, today=TODAY)
    assert "Saving page 2 failed" in res.failed
    by = {c.code: c for c in res.checks}
    assert by["KWD"].ok
    assert by["USD"].saved == "✗ not saved"


def test_a_reply_that_disagrees_is_flagged():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    z.reply_differs = {"KWD": "0.0026"}
    res = roe.apply(z, plan, today=TODAY)
    assert res.checks[0].saved.startswith("✗ shows 0.0026")
    assert not res.checks[0].ok


def test_quick_sale_not_caught_up_yet_is_a_warning_not_a_failure():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    z.qs_lag = {"KWD"}
    res = roe.apply(z, plan, today=TODAY)
    assert res.checks[0].quick_sale.startswith("?")
    assert res.checks[0].ok


def test_the_record_workbook_says_what_happened(tmp_path):
    from openpyxl import load_workbook
    from src import roe_report as rr
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84", "USD": "120.5"})
    res = roe.apply(z, plan, today=TODAY)
    path = tmp_path / "roe.xlsx"
    rr.build_workbook(plan, res, path)
    ws = load_workbook(path)["ROE update"]
    assert ws["A3"].value == "All saved and checked."
    assert [c.value for c in ws[5]][:4] == ["Currency", "Sheet: 1 XXX = BDT",
                                            "Rate before", "Rate saved"]
    assert ws["A6"].value == "KWD" and ws["D6"].value == "0.002501"


# ---- the review's hardening --------------------------------------------------

def test_a_reordered_page_stops_the_save():
    """KWD's box moving to another row must not let its rate land in
    another currency's box."""
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    PAGES[1][2], PAGES[1][3] = PAGES[1][3], PAGES[1][2]       # EUR <-> KWD
    try:
        res = roe.apply(z, plan, today=TODAY)
    finally:
        PAGES[1][2], PAGES[1][3] = PAGES[1][3], PAGES[1][2]
    assert not z.saves
    assert "changed since the preview" in res.failed


def test_an_approval_works_once():
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    page = plan.pages[1]
    form = plan.save_form(page, True)
    permit = plan.permits(page, True)
    cl.check_allowed("POST", BASE + cl.ROE, form, permit)
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed("POST", BASE + cl.ROE, form, permit)


def test_a_save_is_never_retried():
    """A timeout after Zenith got the save must not send it again."""
    class Session:
        def __init__(self):
            self.posts = 0

        def post(self, url, **kw):
            self.posts += 1
            raise ConnectionError("timed out")
    sess = Session()
    client = cl.PricingClient(type("Z", (), {"session": sess})(),
                              base_url=BASE, attempts=3)
    z = FakeRoe()
    plan = plan_for(z, {"KWD": "399.84"})
    page = plan.pages[1]
    with pytest.raises(cl.PricingError):
        client.roe_post(plan.save_form(page, True), permit=plan.permits(page,
                                                                       True))
    assert sess.posts == 1
    assert client.attempts == 3          # reads still retry


@pytest.mark.parametrize("method,path,form", [
    ("GET", cl.ROE + "?hAction=Save&DeviseTaux1=9", None),
    ("POST", cl.ROE + "?hAction=Save", [("hAction", "")]),
    ("POST", cl.ROE, [("hAction", ""), ("haction", "Save")]),
    ("POST", cl.ROE, [("hAction", ""), ("DeviseTaux3", "9"),
                      ("SAVE_DeviseTaux3", "1"), ("DeviseTaux3", "1")]),
])
def test_odd_roe_requests_are_refused(method, path, form):
    with pytest.raises(cl.ReadOnlyViolation):
        cl.check_allowed(method, BASE + path, form)


def test_a_change_over_ten_percent_is_not_saved():
    """40.1 typed for 401: a typo, not a market move."""
    plan = plan_for(FakeRoe(), {"KWD": "40.1", "AED": "33.5"})
    assert [c.code for c in plan.changes] == ["AED"]
    why = dict(plan.skipped)["KWD"]
    assert "NOT saved" in why and "+" in why


def test_a_form_without_the_linked_tax_field_is_not_saved():
    z = FakeRoe()
    plan = plan_for(z, {"USD": "120.5"})
    page = plan.pages[2]
    page.form = [(k, v) for k, v in page.form if k != "UpdateLinkedTaxes"]
    with pytest.raises(roe.RoeError, match="not as recorded"):
        plan.save_form(page, True)


def test_a_click_handler_cannot_change_the_rebuilt_form():
    html = page_html(1, FakeRoe().state).replace(
        "</html>", "<script>function go(){ theForm.hAction.value = 'Save'; "
                   "}</script></html>")
    assert dict(roe.browser_form(html, "frmDevise"))["hAction"] == ""
