"""The Flight Change Authenticator's audit workbook.

Beyond "it writes a file", these pin the things that would quietly mislead a
reader: a blank row swallowed under a header, an unjustified reissue not standing
out, the unknown-marker sheet losing its contents, or the thresholds used not
being recorded alongside the verdicts.
"""

from __future__ import annotations

from datetime import datetime, timedelta

from openpyxl import load_workbook

from src.flight_change_auth import AuthConfig, authenticate_pnr
from src.flight_change_report import (
    VERDICT_COLUMNS,
    build_flight_change_audit_path,
    write_flight_change_audit,
)


class _Agent:
    def __init__(self, name, user_id, department=""):
        self.name, self.user_id, self.department = name, user_id, department


class _Event:
    def __init__(self, ts, desc="", etype="", agent=None, is_reissue=False):
        self.timestamp = ts
        self.raw_description = desc
        self.event_type = etype
        self.agent = agent or _Agent("", "")
        self.is_reissue = is_reissue


_T0 = datetime(2025, 10, 1, 10, 0)
_REISSUE = "Issued->Exchanged | IATA Coupon status :I ->E"


def _fixture():
    """Three PNRs: one unearned, one earned by the same agent who moved it, one
    with an event type no grammar knows."""
    return {
        "08200H": [
            _Event(_T0, "DACSHJ/10/10/2025/ 20:45->00:30/1 ==> 10/10/2025/ 21:25->00:30/1",
                   "Changing flight time", _Agent("Tanni", "tanni7196", "BO-3 Revenue")),
            _Event(_T0 + timedelta(days=1), _REISSUE, "Ticket Modification",
                   _Agent("Akhter", "alaya1751", "BO-1 Central Reservation"), True),
        ],
        "08DS6J": [
            _Event(_T0, "CXBDAC/11/11/2025/ 10:50->11:55 ==> 11/11/2025/ 10:20->11:25",
                   "Changing flight time", _Agent("Robin", "robin5914", "BO-3 Revenue")),
            _Event(_T0 + timedelta(days=2), _REISSUE, "Ticket Modification",
                   _Agent("Robin", "robin5914", "BO-3 Revenue"), True),
        ],
        "08HA2G": [
            _Event(_T0, "something new", "Mystery Type", _Agent("X", "x1")),
            _Event(_T0 + timedelta(hours=3), _REISSUE, "Ticket Modification",
                   _Agent("Y", "y2"), True),
        ],
    }


def _build(tmp_path, cfg=None):
    cfg = cfg or AuthConfig()
    events = _fixture()
    cases = [c for pnr, evs in events.items() for c in authenticate_pnr(pnr, evs, cfg)]
    path = build_flight_change_audit_path(tmp_path)
    write_flight_change_audit(path, cases, cfg, events)
    return load_workbook(path), cases


def test_workbook_has_the_five_sheets(tmp_path):
    wb, _ = _build(tmp_path)
    assert wb.sheetnames == ["Verdicts", "Timeline", "Agents", "Review", "Config"]


def test_no_blank_row_is_left_under_any_header(tmp_path):
    """A header written cell-by-cell used to bump max_row, so append() skipped a
    line and every sheet opened with an empty first row."""
    wb, cases = _build(tmp_path)
    assert wb["Verdicts"].cell(2, 1).value is not None
    assert wb["Agents"].cell(2, 1).value is not None
    assert wb["Config"].cell(2, 1).value is not None
    assert wb["Review"].cell(4, 1).value is not None          # header sits on row 3
    assert wb["Verdicts"].max_row == 1 + len(cases)


def test_verdict_row_carries_the_evidence(tmp_path):
    wb, _ = _build(tmp_path)
    ws = wb["Verdicts"]
    assert [c.value for c in ws[1]] == VERDICT_COLUMNS
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    r = rows["08200H"]
    assert ws.cell(r, 2).value == "BELOW_THRESHOLD"
    assert "60 min" in ws.cell(r, 3).value                     # the reason in words
    assert ws.cell(r, 4).value == "International"
    assert ws.cell(r, 5).value == "DAC-SHJ"
    assert ws.cell(r, 8).value == 40                           # governing shift
    assert ws.cell(r, 11).value == "Tanni" and ws.cell(r, 14).value == "Akhter"


def test_unearned_reissue_is_visually_flagged(tmp_path):
    wb, _ = _build(tmp_path)
    ws = wb["Verdicts"]
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    bad = ws.cell(rows["08200H"], 2).fill.fgColor.rgb
    good = ws.cell(rows["08DS6J"], 2).fill.fgColor.rgb
    assert bad.endswith("FFC7CE") and good.endswith("C6EFCE")


def test_same_agent_case_is_marked(tmp_path):
    wb, _ = _build(tmp_path)
    ws = wb["Verdicts"]
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(rows["08DS6J"], 18).value == "YES"          # moved it AND reissued
    assert ws.cell(rows["08200H"], 18).value in ("", None)
    assert "SAME AGENT" in (ws.cell(rows["08DS6J"], 19).value or "")


def test_unknown_marker_reaches_the_review_sheet(tmp_path):
    wb, _ = _build(tmp_path)
    found = {wb["Review"].cell(r, 1).value: wb["Review"].cell(r, 2).value
             for r in range(4, wb["Review"].max_row + 1)}
    assert found.get("Mystery Type") == 1


def test_timeline_marks_what_each_event_was_recognised_as(tmp_path):
    wb, _ = _build(tmp_path)
    ws = wb["Timeline"]
    kinds = {(ws.cell(r, 1).value, ws.cell(r, 3).value): ws.cell(r, 6).value
             for r in range(2, ws.max_row + 1)}
    assert kinds[("08200H", "Changing flight time")] == "TIME_REVISION"
    assert kinds[("08200H", "Ticket Modification")] == "REISSUE"
    assert kinds[("08HA2G", "Mystery Type")] in ("", None)     # honestly unrecognised


def test_agents_sheet_counts_by_verdict(tmp_path):
    wb, _ = _build(tmp_path)
    ws = wb["Agents"]
    by_login = {ws.cell(r, 2).value: [c.value for c in ws[r]]
                for r in range(2, ws.max_row + 1)}
    assert by_login["alaya1751"][4] == 0 and by_login["alaya1751"][5] == 1   # 1 below
    assert by_login["robin5914"][4] == 1                                     # justified
    assert by_login["robin5914"][10] == 1                                    # also moved it


def test_config_sheet_records_the_thresholds_used(tmp_path):
    cfg = AuthConfig(domestic_minutes=45, international_minutes=90,
                     reissue_window_days=14, inclusive=False)
    wb, _ = _build(tmp_path, cfg)
    got = {wb["Config"].cell(r, 1).value: wb["Config"].cell(r, 2).value
           for r in range(2, wb["Config"].max_row + 1)}
    assert got["Domestic threshold (minutes)"] == 45
    assert got["International threshold (minutes)"] == 90
    assert got["Reissue window (days)"] == 14
    assert got["Exactly at the threshold qualifies"] == "NO"


def test_workbook_can_be_written_with_no_timeline(tmp_path):
    """events_by_pnr is optional — the verdicts must still write."""
    cfg = AuthConfig()
    cases = [c for pnr, evs in _fixture().items()
             for c in authenticate_pnr(pnr, evs, cfg)]
    path = build_flight_change_audit_path(tmp_path)
    write_flight_change_audit(path, cases, cfg)
    wb = load_workbook(path)
    assert wb["Verdicts"].max_row == 1 + len(cases)
    assert wb["Timeline"].max_row == 1
