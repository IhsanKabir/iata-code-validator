"""Every suggested change comes with the evidence behind it."""
from openpyxl import load_workbook

from src import fleet_rotation as fr
from src import route_optimisation as ro
from src import route_optimisation_report as ror
from src import route_reasons as rr


def _view(route, lf=0.97, rask_rev=1_000_000.0):
    return ro.RouteView(route, our_flights=3 / 7, our_seats=31,
                        our_taken=31 * lf, our_legs=40,
                        aircraft="ATR-72-600",
                        rivals=[ro.Rival("6E", 1, 180, "Airbus A320"),
                                ro.Rival("BG", 1, 74, "DH8")],
                        distance_km=330.0, rask_revenue=rask_rev,
                        rask_seats=2_000.0)


def _kolkata():
    return ro.Result(routes=[_view("DAC-CCU"), _view("CCU-DAC")])


def _check(good=("Mon", "Wed", "Thu", "Sat"), flown=("Tue", "Fri", "Sun")):
    c = fr.RotationCheck(pair="DAC ⇄ CCU", base="DAC", fam="ATR 72",
                         rotation=270, flown_days=flown)
    for wd in fr.WEEKDAYS:
        c.by_weekday[wd] = ((13, 0, 13, 16 * 60 + 25, 2) if wd in good
                            else (0, 0, 13, None, 0))
    return c


def test_a_full_out_flown_pair_gets_a_suggestion_with_four_reasons():
    got = rr.advise_all(_kolkata(), [_check()])
    assert len(got) == 1
    item = got[0]
    assert item.pair == "DAC ⇄ CCU" and item.kind == "add"
    heads = [r.split(":")[0] for r in item.reasons]
    assert heads == ["Demand", "Competition", "Value", "Aircraft"]


def test_the_reasons_name_the_rivals_their_aircraft_and_our_share():
    text = " ".join(rr.advise_all(_kolkata(), [_check()])[0].reasons)
    assert "6E daily Airbus A320" in text and "BG daily DH8" in text
    assert "3/week" in text
    assert "of the seats" in text and "of the departures" in text


def test_going_daily_is_suggested_only_when_every_missing_day_has_a_slot():
    item = rr.advise_all(_kolkata(), [_check()])[0]
    assert item.action == "Go daily — add Monday, Wednesday, Thursday " \
                          "and Saturday"
    assert "16:25" in item.reasons[-1] and "2 aircraft" in item.reasons[-1]


def test_some_missing_days_with_a_slot_are_named():
    item = rr.advise_all(_kolkata(), [_check(good=("Mon",))])[0]
    assert item.action == "Add a flight on Monday"


def test_no_slot_says_an_aircraft_is_needed():
    item = rr.advise_all(_kolkata(), [_check(good=())])[0]
    assert "needs an extra or bigger aircraft" in item.action
    assert "free on only 0% of days" in item.reasons[-1]


def test_no_slot_but_an_unused_aircraft_says_so_and_names_the_standby():
    """Two ATRs the timetable never needs are on the ground most days.
    Saying 'needs an extra aircraft' would be wrong; using them is a
    planning call, so the suggestion says whose."""
    check = _check(good=())
    for wd in fr.WEEKDAYS:
        check.by_weekday[wd] = (6, 13, 13, 5 * 60, 1)
    item = rr.advise_all(_kolkata(), [check])[0]
    assert "standby" in item.action and "planning" in item.action
    assert "free on only 46% of days" in item.reasons[-1]


def test_a_missing_distance_is_said_not_skipped():
    res = ro.Result(routes=[_view("DAC-CCU", rask_rev=None),
                            _view("CCU-DAC", rask_rev=None)])
    for v in res.routes:
        v.distance_km = None
    text = " ".join(rr.advise_all(res, [_check()])[0].reasons)
    assert "distance is missing" in text


def test_only_the_full_direction_carries_the_demand_case():
    res = ro.Result(routes=[_view("DAC-CCU"), _view("CCU-DAC", lf=0.70)])
    demand = rr.advise_all(res, [_check()])[0].reasons[0]
    assert "DAC-CCU is where seats run out" in demand


def test_a_half_empty_pair_is_flagged_for_review():
    res = ro.Result(routes=[_view("DAC-CCU", lf=0.5),
                            _view("CCU-DAC", lf=0.5)])
    item = rr.advise_all(res)[0]
    assert (item.kind, item.action) == ("review", "Review capacity")


def test_a_steady_pair_gets_no_suggestion():
    res = ro.Result(routes=[_view("DAC-CCU", lf=0.8),
                            _view("CCU-DAC", lf=0.8)])
    assert rr.advise_all(res) == []


def test_the_summary_sheet_explains_each_suggestion(tmp_path):
    out = tmp_path / "ro.xlsx"
    ror.build_workbook(_kolkata(), out, None, [_check()])
    ws = load_workbook(out)["Summary"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row
                    if c.value is not None)
    assert "WHAT TO CHANGE, AND WHY" in text
    assert "Go daily" in text and "• Demand:" in text
