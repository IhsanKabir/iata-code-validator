"""Audit workbook for the Flight Change Authenticator.

Sheets:
  Verdicts   one row per reissue, colour-coded, with the reason in words
  Timeline   every history event per PNR, marked with what it was recognised as —
             the evidence trail behind each verdict
  Agents     who reissued, and how often it was not justified
  Review     event types no grammar recognises (the discovery loop for other ways
             a flight change gets recorded)
  Config     the thresholds used, so a result can be reproduced

Every verdict is decision support, not a finding. The Reason column and the
Timeline exist so a human can check a case before anyone is accused of anything.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .flight_change_auth import changes_in_event, suspicious_flags

VERDICT_COLUMNS = [
    "PNR", "Verdict", "Reason", "Sector", "Route",
    "Original Departure", "Revised Departure", "Governing Shift (min)",
    "Cumulative Shift (min)", "Change At", "Changed By", "Changed By (login)",
    "Reissue At", "Reissued By", "Reissued By (login)", "Dept",
    "Days To Reissue", "Same Agent", "Observations",
]

_VERDICT_FILL = {
    "JUSTIFIED": "C6EFCE",        # green — earned
    "BELOW_THRESHOLD": "FFC7CE",  # red — a free reissue that was not earned
    "OUTSIDE_WINDOW": "FFC7CE",
    "NO_CHANGE_FOUND": "FFC7CE",
    "NEEDS_REVIEW": "FFEB9C",     # amber — unknown marker, must be looked at
}

_NAVY = PatternFill("solid", fgColor="1F3864")
_WHITE_BOLD = Font(bold=True, color="FFFFFF", size=10)


def _head(ws, cols, row: int = 1) -> None:
    for j, h in enumerate(cols, 1):
        c = ws.cell(row, j, h)
        c.font = _WHITE_BOLD
        c.fill = _NAVY
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # NB: ws.cell(...) CREATES the cell and bumps max_row, which made append()
    # skip a line and leave a blank row under every header. Build the coordinate.
    ws.freeze_panes = f"A{row + 1}"


def _dt(v) -> str:
    return v.strftime("%d/%m/%Y %H:%M") if v else ""


def write_flight_change_audit(path: Path, cases, cfg, events_by_pnr=None) -> Path:
    wb = Workbook()

    # --- Verdicts -----------------------------------------------------------
    ws = wb.active
    ws.title = "Verdicts"
    _head(ws, VERDICT_COLUMNS)
    r = 2
    for c in cases:
        gov = c.governing_change
        ws.append([
            c.pnr, c.verdict, c.reason, c.sector, c.route,
            _dt(getattr(gov, "original_dep", None)),
            _dt(getattr(gov, "revised_dep", None)),
            getattr(gov, "shift_minutes", None),
            c.cumulative_shift,
            _dt(c.change_at), c.changed_by, c.changed_by_login,
            _dt(c.reissue_at), c.reissued_by, c.reissued_by_login, c.reissued_by_dept,
            round(c.days_to_reissue, 2) if c.days_to_reissue is not None else None,
            "YES" if c.same_agent else "",
            "; ".join(suspicious_flags(c, cfg)),
        ])
        fill = _VERDICT_FILL.get(c.verdict)
        if fill:
            ws.cell(r, 2).fill = PatternFill("solid", fgColor=fill)
        if c.same_agent:
            ws.cell(r, 18).fill = PatternFill("solid", fgColor="FFC7CE")
        r += 1
    for col, w in (("A", 11), ("B", 17), ("C", 46), ("D", 14), ("E", 11),
                   ("F", 17), ("G", 17), ("H", 16), ("I", 16), ("J", 17),
                   ("K", 20), ("L", 16), ("M", 17), ("N", 20), ("O", 16),
                   ("P", 24), ("Q", 14), ("R", 12), ("S", 52)):
        ws.column_dimensions[col].width = w

    # --- Timeline -----------------------------------------------------------
    ws = wb.create_sheet("Timeline")
    _head(ws, ["PNR", "Time", "Type", "Agent", "Login", "Recognised As", "Description"])
    for pnr, events in (events_by_pnr or {}).items():
        for e in sorted(events, key=lambda x: (getattr(x, "timestamp", None)
                                               or datetime.min)):
            found = changes_in_event(e, cfg)
            if getattr(e, "is_reissue", False):
                kind = "REISSUE"
            elif found:
                kind = "/".join(sorted({f.kind for f in found}))
            else:
                kind = ""
            agent = getattr(e, "agent", None)
            ws.append([
                pnr, _dt(getattr(e, "timestamp", None)),
                getattr(e, "event_type", ""),
                getattr(agent, "name", "") if agent else "",
                getattr(agent, "user_id", "") if agent else "",
                kind,
                (getattr(e, "raw_description", "") or "")[:400],
            ])
    for col, w in (("A", 11), ("B", 17), ("C", 26), ("D", 22), ("E", 14),
                   ("F", 18), ("G", 90)):
        ws.column_dimensions[col].width = w

    # --- Agents -------------------------------------------------------------
    ws = wb.create_sheet("Agents")
    _head(ws, ["Agent", "Login", "Dept", "Reissues", "Justified", "Below Threshold",
               "Outside Window", "No Change", "Needs Review", "% Not Justified",
               "Also Moved The Schedule"])
    per: dict = defaultdict(Counter)
    meta: dict = {}
    for c in cases:
        key = c.reissued_by_login or c.reissued_by or "(unknown)"
        per[key][c.verdict] += 1
        per[key]["total"] += 1
        if c.same_agent:
            per[key]["same_agent"] += 1
        meta.setdefault(key, (c.reissued_by, c.reissued_by_dept))
    for key, cnt in sorted(per.items(), key=lambda kv: -kv[1]["total"]):
        name, dept = meta.get(key, ("", ""))
        total = cnt["total"]
        bad = total - cnt["JUSTIFIED"]
        ws.append([name, key, dept, total, cnt["JUSTIFIED"], cnt["BELOW_THRESHOLD"],
                   cnt["OUTSIDE_WINDOW"], cnt["NO_CHANGE_FOUND"], cnt["NEEDS_REVIEW"],
                   (bad / total) if total else None, cnt["same_agent"]])
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        row[0].number_format = "0%"
    for col, w in (("A", 24), ("B", 14), ("C", 26), ("J", 15), ("K", 22)):
        ws.column_dimensions[col].width = w

    # --- Review (the discovery loop) ----------------------------------------
    ws = wb.create_sheet("Review")
    note = ws.cell(1, 1, "Event types seen on reissued PNRs that NO grammar "
                         "recognises. A new way of recording a flight change shows "
                         "up here — add its grammar and re-run.")
    note.font = Font(italic=True, color="595959", size=9)
    _head(ws, ["Unrecognised Event Type", "Reissues Affected"], row=3)
    seen: Counter = Counter()
    for c in cases:
        for t in c.unclassified_types:
            seen[t] += 1
    for t, n in seen.most_common():
        ws.append([t, n])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    # --- Config -------------------------------------------------------------
    ws = wb.create_sheet("Config")
    _head(ws, ["Setting", "Value"])
    for k, v in (
        ("Domestic threshold (minutes)", cfg.domestic_minutes),
        ("International threshold (minutes)", cfg.international_minutes),
        ("Reissue window (days)", cfg.reissue_window_days),
        ("Exactly at the threshold qualifies", "YES" if cfg.inclusive else "NO"),
        ("Judged on", cfg.basis),
        ("Domestic airports", ", ".join(sorted(cfg.bd_airports))),
        ("Generated", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ):
        ws.append([k, v])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 44

    wb.save(path)
    return Path(path)


def build_flight_change_audit_path(folder: Path) -> Path:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"flight_change_audit_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
