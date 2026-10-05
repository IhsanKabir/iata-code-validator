"""The season check workbook.

Sheets:
  Summary       which seasons, which routes, how much was compared, how
                many differences each part found, and what could not be read
  Fares         Fare Editor rows that differ, are missing or are new
  Sales check   Fare Editor adult fares against what Quick Sale sells
  QS taxes      Quick Sale's total tax per fare, season against season
  Taxes         the tax screen's list per route (YQ/YR included)
  Penalties     change / refund / no-show fees per class, route and window
  Not checked   anything the check could not read -- never silently dropped

An empty part sheet says so in words: "no differences" must be visibly a
result, not a blank page that might be a failure.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook

from .counter_master import BAD, GOOD, GREY, NAVY, PAPER, WARN, _cell, _headers

PARTS = ("Fares", "Sales check", "QS taxes", "Taxes", "Penalties")
HEADERS = ("Route", "Item", "Field / window", "Finding", "Season A",
           "Value A", "Season B", "Value B")
WIDTHS = (26, 44, 30, 20, 18, 22, 18, 22)
_TONE = {"different": WARN, "amount differs": WARN, "missing": BAD,
         "set up, not on sale": BAD, "new": PAPER, "duplicate": BAD}


def _widths(ws, widths) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.sheet_view.showGridLines = False


def _title(ws, text: str, note: str) -> int:
    _cell(ws, 1, 1, text, bold=True, size=15, color=NAVY)
    _cell(ws, 2, 1, note, size=9, color=GREY)
    return 4


def write_summary(ws, res) -> None:
    _widths(ws, (28, 16, 16, 16, 16, 16, 16))
    r = _title(ws, "Season check — read only",
               f"Run {res.started} to {res.finished}. Nothing in Zenith was "
               f"changed: only searches and displays were read.")
    for label, value in (("Base season", res.base),
                         ("Compared with", ", ".join(res.others)),
                         ("Routes", f"{len(res.routes):,}: "
                                    + ", ".join(res.routes[:40])
                                    + (" …" if len(res.routes) > 40 else ""))):
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, value)
        r += 1
    r += 1
    r = _headers(ws, r, ("Part", "Compared", "Differences", "Missing",
                         "New", "Other", "Result"))
    for part in PARTS:
        counts = res.count(part)
        diff = counts.get("different", 0) + counts.get("amount differs", 0)
        missing = counts.get("missing", 0) + counts.get(
            "set up, not on sale", 0)
        new = counts.get("new", 0)
        other = sum(counts.values()) - diff - missing - new
        compared = res.compared.get(part, 0)
        total = sum(counts.values())
        verdict = ("not run" if not compared and not total else
                   "all match" if not total else f"{total:,} to look at")
        tone = PAPER if verdict == "not run" else GOOD if not total else BAD
        for j, v in enumerate((part, compared, diff, missing, new, other,
                               verdict), start=1):
            _cell(ws, r, j, v, fill=tone if j == 7 else None, border=True,
                  fmt="#,##0" if isinstance(v, int) else None)
        r += 1
    r += 1
    if res.problems:
        _cell(ws, r, 1, f"{len(res.problems):,} item(s) could not be "
                        f"checked — see 'Not checked'.", bold=True,
              color="B00020")
    else:
        _cell(ws, r, 1, "Everything asked for was read.", color=GREY)
    ws.freeze_panes = None


def write_part(ws, res, part: str) -> int:
    _widths(ws, WIDTHS)
    rows = [f for f in res.findings if f.part == part]
    r = _title(ws, part, _NOTES[part])
    if not rows:
        compared = res.compared.get(part, 0)
        _cell(ws, r, 1, f"No differences — {compared:,} compared."
              if compared else "This part was not run.",
              bold=True, fill=GOOD if compared else PAPER)
        return 0
    start = r
    r = _headers(ws, r, HEADERS)
    for f in sorted(rows, key=lambda f: (f.route, f.item, f.detail)):
        for j, v in enumerate((f.route, f.item, f.detail, f.kind, f.season_a,
                               f.value_a, f.season_b, f.value_b), start=1):
            _cell(ws, r, j, v, border=True,
                  fill=_TONE.get(f.kind) if j == 4 else None)
        r += 1
    ws.auto_filter.ref = f"A{start}:H{r - 1}"
    ws.freeze_panes = f"A{start + 1}"
    return len(rows)


_NOTES = {
    "Fares": "Fare Editor rows (fare basis × route × passenger type); every "
             "field compared, currency included.",
    "Sales check": "Adult fares as set in the Fare Editor against Quick "
                   "Sale's display in the same currency, per season.",
    "QS taxes": "Quick Sale total minus fare (with taxes − without) for each "
                "fare basis, season against season.",
    "Taxes": "The tax screen's list for each route on a day in each season "
             "— codes, values, YQ/YR.",
    "Penalties": "Change and refund fees per fare class, route, passenger "
                 "type and time window.",
}


def write_problems(ws, res) -> None:
    _widths(ws, (120,))
    r = _title(ws, "Not checked", "What could not be read. These are NOT "
                                  "passes — check them by hand or re-run.")
    if not res.problems:
        _cell(ws, r, 1, "Nothing — every item asked for was read.",
              fill=GOOD)
        return
    for p in res.problems:
        _cell(ws, r, 1, p, wrap=True)
        r += 1


def build_workbook(res, out_path) -> None:
    wb = Workbook()
    wb.active.title = "Summary"
    write_summary(wb["Summary"], res)
    for part in PARTS:
        write_part(wb.create_sheet(part), res, part)
    write_problems(wb.create_sheet("Not checked"), res)
    wb.save(str(out_path))


def default_filename() -> str:
    return f"Season_Check_{date.today():%Y-%m-%d}.xlsx"
