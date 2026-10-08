"""The ROE update record: what was asked, what was saved, how it checked.

One workbook per run, kept as the day's evidence: the sheet's rate, the
rate saved, the rate before, and the five checks for each currency.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .counter_master import BAD, GOOD, GREY, NAVY, PAPER, WARN, _cell, _headers
from .roe import rate_text

HEADERS = ("Currency", "Sheet: 1 XXX = BDT", "Rate before", "Rate saved",
           "Change %", "Linked taxes", "1. Zenith reply", "2. Fresh read",
           "3. Last modified", "4. History", "5. Quick Sale")
WIDTHS = (10, 18, 13, 13, 10, 13, 18, 16, 24, 34, 30)


def _tone(text: str):
    return GOOD if text.startswith("✓") else BAD if text.startswith(
        "✗") else WARN if text.startswith("?") else None


def build_workbook(plan, res, out_path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ROE update"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _cell(ws, 1, 1, "Exchange rates (ROE) — update record", bold=True,
          size=15, color=NAVY)
    sheet_day = f"{plan.sheet_day:%d %b %Y}" if plan.sheet_day else "unknown"
    _cell(ws, 2, 1, f"Sheet dated {sheet_day} · saved {res.started} → "
                    f"{res.finished} · linked taxes "
                    f"{'updated' if res.update_linked else 'not updated'}",
          size=9, color=GREY)
    verdict = ("All saved and checked." if not res.failed and res.untouched_ok
               and all(c.ok for c in res.checks)
               else res.failed or "Saved — but see the checks marked ✗.")
    _cell(ws, 3, 1, verdict, bold=True,
          fill=GOOD if verdict.startswith("All") else BAD)
    r = _headers(ws, 5, HEADERS)
    by_code = {c.code: c for c in plan.changes}
    for chk in res.checks:
        c = by_code[chk.code]
        vals = (chk.code, float(c.bdt_per_unit), chk.old, chk.new,
                round(float(c.pct), 3), ", ".join(c.linked_taxes),
                chk.saved, chk.reread, chk.modified, chk.history,
                chk.quick_sale)
        for j, v in enumerate(vals, start=1):
            _cell(ws, r, j, v, border=True,
                  fill=_tone(v) if j >= 7 and isinstance(v, str) else
                  WARN if j == 5 and c.big else None)
        r += 1
    r += 1
    _cell(ws, r, 1, "Every other currency: " + (
        "unchanged, as expected." if res.untouched_ok
        else " · ".join(res.untouched_notes)), bold=True,
        fill=GOOD if res.untouched_ok else BAD)
    r += 2
    if plan.same:
        _cell(ws, r, 1, "Already at the sheet's rate: "
                        + ", ".join(plan.same), color=GREY)
        r += 1
    for code, why in plan.skipped:
        _cell(ws, r, 1, f"{code}: {why}", color=GREY)
        r += 1
    for w in plan.warnings:
        _cell(ws, r, 1, f"Note: {w}", fill=PAPER)
        r += 1
    ws.freeze_panes = "B6"
    wb.save(str(out_path))


def default_filename() -> str:
    return f"ROE_Update_{date.today():%Y-%m-%d}.xlsx"


def preview_rows(plan) -> list:
    """Rows for the on-screen table: (code, sheet, now, new, change, note)."""
    rows = []
    for c in plan.changes:
        rows.append((c.code, f"{c.bdt_per_unit:f}", rate_text(c.old),
                     rate_text(c.new), f"{c.pct:+.2f}%",
                     ("linked: " + ", ".join(c.linked_taxes)
                      if c.linked_taxes else "")
                     + (" · BIG CHANGE" if c.big else "")
                     + (f" · sheet note {c.note}" if c.note else "")))
    for code in plan.same:
        rows.append((code, "", "", "", "same", "already at the sheet's rate"))
    for code, why in plan.skipped:
        rows.append((code, "", "", "", "—", why))
    return rows
