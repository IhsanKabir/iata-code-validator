"""Every FFP member, out of a search that lists 50 at a time.

Zenith's FFP search shows at most 50 rows but always says how many matched
(see zenith_ffp). So the whole membership is reached by narrowing: search
FFP numbers beginning 0, 1, ... 9; any beginning that matches more than 50
is split into ten longer ones (10, 11, ... 19), and so on, until every
search lists all it matched. Every member is found this way whether the
field matches from the start of the number or anywhere in it: a member's
own number contains each of its own beginnings, so the narrowing follows
that path down to a search small enough to list it.

Checks come first, because the whole method rests on them:

* each level is searched on its own, so the counts Zenith holds per level
  are known, and the finished list can be checked against them;
* narrowing needs a field that matches PART of a value. On the live system
  the FFP number field matches whole numbers only -- a real number less its
  last digit found nothing. So each way is tried on a real member, and the
  first that brings that member back is used: part of the FFP number as
  typed, then with a '%' or '*' wildcard; then part of the last name within
  one level, the same three ways, noting whether case matters. If none
  works, the run stops saying so rather than returning a partial list.

Narrowing by last name works like narrowing by number, a level at a time:
names beginning A-Z, then each beginning that matches more than 50 grows
by one more character.

Progress is kept in a small local database as it goes -- each search that
finished, each member found -- so Stop, a crash or a signed-out session
loses nothing, and the next run carries on from where this one stopped.
It holds customers' personal details: it stays on this machine.
"""
from __future__ import annotations

import json
import sqlite3
import threading
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path

from . import zenith_ffp as zf

DIGITS = "0123456789"
#: Deeper than any customer number; a beginning this long that still
#: matches more than 50 is recorded as a gap rather than split forever.
MAX_DIGITS = 12

PENDING, DONE, SPLIT, CAPPED = "pending", "done", "split", "capped"
#: Taken by a worker and not yet answered. A run that stops leaves these
#: behind; the next run hands them out again.
RUNNING = "running"

#: The order levels are collected in: the small, valuable ones first.
LEVEL_ORDER = ("Gold", "Platinum", "Titanium", "Silver")

_MEMBER_COLS = ("ffp_number", "customer_id", "last_name", "first_name",
                "birth_date", "email", "phone", "id_number", "level",
                "miles")


def _level_of(prefix: str) -> str:
    return prefix.split("|", 1)[0] if "|" in prefix else ""


def _rank(prefix: str) -> int:
    level = _level_of(prefix)
    return LEVEL_ORDER.index(level) if level in LEVEL_ORDER else         len(LEVEL_ORDER)


class PartialMatchUnsupported(Exception):
    """The FFP number field only matches whole numbers."""


class FFPStore:
    """The run's progress and every member found, on this machine."""

    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # one connection, several workers: every statement goes through this
        self._lock = threading.RLock()
        self.con = sqlite3.connect(str(self.path), check_same_thread=False)
        # A commit per search; write-ahead logging keeps each one cheap and
        # still survives the app being closed or crashing mid-run.
        self.con.execute("PRAGMA journal_mode=WAL")
        self.con.execute("PRAGMA synchronous=NORMAL")
        cols = ", ".join(f"{c} TEXT NOT NULL DEFAULT ''"
                         for c in _MEMBER_COLS[1:])
        self.con.executescript(f"""
            CREATE TABLE IF NOT EXISTS members (
                ffp_number TEXT PRIMARY KEY, {cols},
                found_at TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS frontier (
                prefix TEXT PRIMARY KEY, state TEXT NOT NULL,
                total INTEGER NOT NULL DEFAULT -1);
            CREATE TABLE IF NOT EXISTS meta (
                key TEXT PRIMARY KEY, value TEXT NOT NULL);
        """)
        # rank (level order) and depth are kept as columns, indexed, so
        # picking the next search does not scan a queue that runs to
        # tens of thousands; a database from an earlier version gets them
        have = {r[1] for r in self.con.execute("PRAGMA table_info(frontier)")}
        if "rank" not in have:
            self.con.execute("ALTER TABLE frontier ADD COLUMN rank INTEGER "
                             "NOT NULL DEFAULT 0")
            self.con.execute("ALTER TABLE frontier ADD COLUMN depth INTEGER "
                             "NOT NULL DEFAULT 0")
            for (prefix,) in self.con.execute(
                    "SELECT prefix FROM frontier").fetchall():
                self.con.execute(
                    "UPDATE frontier SET rank=?, depth=? WHERE prefix=?",
                    (_rank(prefix), len(prefix), prefix))
        self.con.execute("CREATE INDEX IF NOT EXISTS frontier_next ON "
                         "frontier (state, rank, depth, prefix)")
        self.con.commit()

    def close(self) -> None:
        self.con.close()

    # ---- meta ------------------------------------------------------------
    def get(self, key: str, default: str = "") -> str:
        with self._lock:
            row = self.con.execute("SELECT value FROM meta WHERE key=?",
                                   (key,)).fetchone()
        return row[0] if row else default

    def put(self, key: str, value) -> None:
        with self._lock:
            self.con.execute("INSERT OR REPLACE INTO meta VALUES (?, ?)",
                             (key, str(value)))
            self.con.commit()

    def level_totals(self) -> dict:
        return {lv: int(self.get(f"level_total:{lv}", "-1"))
                for lv in zf.LEVELS}

    # ---- members ---------------------------------------------------------
    def save_members(self, members) -> int:
        now = time.strftime("%Y-%m-%d %H:%M:%S")
        rows = [tuple(asdict(m)[c] for c in _MEMBER_COLS) + (now,)
                for m in members if m.ffp_number]
        with self._lock:
            before = self.member_count()
            self.con.executemany(
                f"INSERT OR REPLACE INTO members "
                f"({', '.join(_MEMBER_COLS)}, found_at) "
                f"VALUES ({', '.join('?' * (len(_MEMBER_COLS) + 1))})", rows)
            self.con.commit()
            return self.member_count() - before

    def member_count(self) -> int:
        with self._lock:
            return self.con.execute(
                "SELECT count(*) FROM members").fetchone()[0]

    def counts_by_level(self) -> dict:
        with self._lock:
            return dict(self.con.execute(
                "SELECT level, count(*) FROM members GROUP BY level"))

    def members(self):
        with self._lock:
            rows = self.con.execute(
                f"SELECT {', '.join(_MEMBER_COLS)} FROM members "
                f"ORDER BY level, ffp_number").fetchall()
        for row in rows:
            yield dict(zip(_MEMBER_COLS, row))

    # ---- frontier --------------------------------------------------------
    def seed(self, prefixes) -> None:
        with self._lock:
            self.con.executemany(
                "INSERT OR IGNORE INTO frontier (prefix, state, rank, depth) "
                "VALUES (?, ?, ?, ?)",
                [(p, PENDING, _rank(p), len(p)) for p in prefixes])
            self.con.commit()

    def mark(self, prefix: str, state: str, total: int) -> None:
        with self._lock:
            self.con.execute(
                "UPDATE frontier SET state=?, total=? WHERE prefix=?",
                (state, total, prefix))
            self.con.commit()

    @staticmethod
    def _level_filter(levels) -> tuple:
        """SQL keeping name searches of `levels`; number searches carry no
        level (rank past the last level) and are always kept."""
        if levels is None:
            return "", ()
        ranks = sorted({LEVEL_ORDER.index(lv) for lv in levels
                        if lv in LEVEL_ORDER} | {len(LEVEL_ORDER)})
        return (f" AND rank IN ({', '.join('?' * len(ranks))})",
                tuple(ranks))

    def next_pending(self, levels=None):
        where, args = self._level_filter(levels)
        with self._lock:
            row = self.con.execute(
                "SELECT prefix FROM frontier WHERE state=?" + where
                + " ORDER BY rank, depth, prefix LIMIT 1",
                (PENDING, *args)).fetchone()
        return row[0] if row else None

    def claim_next(self, levels=None):
        """Hand one pending search to a worker, so no two take the same."""
        with self._lock:
            prefix = self.next_pending(levels)
            if prefix is not None:
                self.con.execute("UPDATE frontier SET state=? WHERE prefix=?",
                                 (RUNNING, prefix))
                self.con.commit()
            return prefix

    def close_level(self, level: str) -> None:
        """Every queued search of `level` is answered by a single one."""
        with self._lock:
            self.con.execute(
                "UPDATE frontier SET state=? WHERE rank=? AND state IN (?, ?)",
                (DONE, LEVEL_ORDER.index(level), PENDING, RUNNING))
            self.con.commit()

    def release_running(self) -> None:
        """Searches a stopped run left unanswered go back in the queue."""
        with self._lock:
            self.con.execute("UPDATE frontier SET state=? WHERE state=?",
                             (PENDING, RUNNING))
            self.con.commit()

    def running(self) -> int:
        with self._lock:
            return self.con.execute(
                "SELECT count(*) FROM frontier WHERE state=?",
                (RUNNING,)).fetchone()[0]

    def frontier_counts(self) -> dict:
        with self._lock:
            return dict(self.con.execute(
                "SELECT state, count(*) FROM frontier GROUP BY state"))

    def capped(self) -> list:
        with self._lock:
            return [r[0] for r in self.con.execute(
                "SELECT prefix FROM frontier WHERE state=?", (CAPPED,))]

    def reset(self) -> None:
        with self._lock:
            self.con.executescript(
                "DELETE FROM members; DELETE FROM frontier; DELETE FROM meta;")
            self.con.commit()


@dataclass
class Progress:
    members: int = 0
    searches: int = 0
    pending: int = 0
    done: int = 0
    last: str = ""
    level_totals: dict = field(default_factory=dict)
    level_found: dict = field(default_factory=dict)

    @property
    def expected(self) -> int:
        return sum(v for v in self.level_totals.values() if v > 0)


def progress_of(store: FFPStore, searches: int = 0, last: str = "") -> Progress:
    fc = store.frontier_counts()
    return Progress(members=store.member_count(), searches=searches,
                    pending=fc.get(PENDING, 0) + fc.get(RUNNING, 0),
                    done=fc.get(DONE, 0) + fc.get(SPLIT, 0) + fc.get(CAPPED, 0),
                    last=last, level_totals=store.level_totals(),
                    level_found=store.counts_by_level())


#: Ways a search field might accept "the start of" a value: as typed, or
#: with a SQL or shell wildcard after it. Tried in this order.
WILDCARDS = ("", "%", "*")

#: How a last name may continue after its first letter.
_NAME_TAIL = "-.'"
_UPPER = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
_LOWER = _UPPER.lower()
MAX_NAME = 30

NUMBER, NAME = "number", "name"


def _found(page, member) -> bool:
    """The probe brought the sample member back."""
    return any(m.ffp_number == member.ffp_number for m in page.members)


def _check_levels(searcher, store: FFPStore, log) -> dict:
    """Count each level in Zenith; keep the members each search listed."""
    samples = {}
    for level in zf.LEVELS:
        page = searcher.search(level=level)
        store.put(f"level_total:{level}", page.total)
        store.save_members(page.members)
        log(f"{level}: {page.total:,} member(s) in Zenith.")
        if page.members:
            samples[level] = page.members
    if not samples:
        raise PartialMatchUnsupported(
            "No level returned any member, so there is nothing to collect "
            "-- or the level names have changed.")
    return samples


def _try_number(searcher, sample, log):
    """Does the start of an FFP number find its member, with any wildcard?"""
    for wild in WILDCARDS:
        probe = searcher.search(ffp_number=sample.ffp_number[:-1] + wild)
        if _found(probe, sample):
            return wild
    log("Part of an FFP number finds nothing, with or without a wildcard.")
    return None


def _try_name(searcher, samples: dict, log):
    """Does the start of a last name, within a level, find its member?

    Returns (wildcard, case-sensitive) or None. The probe drops the last
    letter of a real member's last name, so a field that matches whole
    names only comes back without that member.
    """
    for level, members in samples.items():
        sample = next((m for m in members if len(m.last_name) >= 3), None)
        if sample is None:
            continue
        part = sample.last_name[:-1]
        for wild in WILDCARDS:
            probe = searcher.search(level=level, last_name=part + wild)
            if not _found(probe, sample):
                continue
            swapped = part.swapcase()
            if swapped == part:
                return wild, False
            other = searcher.search(level=level, last_name=swapped + wild)
            return wild, not _found(other, sample)
        break
    log("Part of a last name finds nothing, with or without a wildcard.")
    return None


def _choose_method(searcher, store: FFPStore, log) -> None:
    samples = _check_levels(searcher, store, log)
    first = next(iter(samples.values()))[0]
    wild = _try_number(searcher, first, log)
    if wild is not None:
        store.put("mode", NUMBER)
        store.put("wild", wild)
        store.seed(DIGITS)
        log("Narrowing by FFP number" + (f" (with '{wild}')" if wild else "")
            + ".")
        return
    got = _try_name(searcher, samples, log)
    if got is None:
        raise PartialMatchUnsupported(
            "Zenith's FFP search matches only whole FFP numbers and whole "
            "last names -- with or without a wildcard -- so the list cannot "
            "be narrowed to fit its 50-row limit. The level counts are "
            "exact, and the members found so far can be exported; the full "
            "list is not collectable through this search.")
    wild, case_sensitive = got
    store.put("mode", NAME)
    store.put("wild", wild)
    store.put("case", "1" if case_sensitive else "0")
    # Where case matters, a name typed in lower case ("afrin") begins with a
    # lower-case letter and is only reached from one.
    first = _UPPER + (_LOWER if case_sensitive else "")
    totals = store.level_totals()
    for level in zf.LEVELS:
        if totals.get(level, 0) > zf.PAGE_CAP:
            store.seed(f"{level}|{c}" for c in first)
    log("Narrowing by last name within each level"
        + (f" (with '{wild}')" if wild else "")
        + (", case-sensitive" if case_sensitive else "") + ".")


def _children(store: FFPStore, prefix: str):
    """The searches one character longer than an over-full one.

    Never one ending in a space: Zenith trims it, so "A " matched exactly
    what "A" did (219), was split again into "A  ", and so on for up to
    30 spaces -- most of a live run's searches went down such chains and
    found nothing new. A space only narrows with a letter after it
    ("Al Amin"), so it is added together with one.
    """
    if store.get("mode") == NUMBER:
        return [prefix + d for d in DIGITS]
    letters = _LOWER + (_UPPER if store.get("case") == "1" else "")
    out = [prefix + c for c in letters + _NAME_TAIL]
    if not prefix.endswith(" "):
        out += [prefix + " " + c for c in letters]
    return out


def repair_space_chains(store: FFPStore) -> int:
    """Clear what the space bug left in a saved run.

    Queued searches whose name part ends in a space, or holds two in a
    row, are dropped; each over-full name start that was split gets its
    "space + letter" searches instead. Returns how many were dropped.
    """
    with store._lock:
        rows = store.con.execute(
            "SELECT prefix, state FROM frontier WHERE instr(prefix, '|') > 0"
        ).fetchall()
        junk = [p for p, st in rows
                if st in (PENDING, RUNNING)
                and (p.endswith(" ") or "  " in p.split("|", 1)[1])]
        store.con.executemany("DELETE FROM frontier WHERE prefix=?",
                              [(p,) for p in junk])
        store.con.commit()
    split = [p for p, st in rows if st == SPLIT and not p.endswith(" ")]
    for p in split:
        store.seed(c for c in _children(store, p)
                   if not c.endswith(" "))
    store.put("space_fix", "1")
    return len(junk)


def _page_extra(store: FFPStore) -> dict:
    """The larger-page setting a test found, as search() arguments."""
    raw = store.get("page_param")
    if not raw:
        return {}
    got = json.loads(raw)
    return {f"extra_{got['where']}": {got["name"]: str(got["value"])}}


def _search(searcher, store: FFPStore, prefix: str):
    wild = store.get("wild")
    extra = _page_extra(store)
    if store.get("mode") == NUMBER:
        return searcher.search(ffp_number=prefix + wild, **extra)
    level, part = prefix.split("|", 1)
    return searcher.search(level=level, last_name=part + wild, **extra)


# ---- does Zenith list more than 50 if asked? -------------------------------

#: Names a "how many rows" setting commonly goes by. NbReponse is the one
#: Zenith's own flight-load report takes.
LIMIT_NAMES = ("PageSize", "pageSize", "MaxResults", "maxResults",
               "NbReponse", "NbResults", "MaxRows", "maxRows", "Take",
               "take", "Top", "top", "Limit", "limit", "ResultsPerPage",
               "NumberOfResults", "RowCount", "Count")
#: Asked for per search. Not more: a page of 5,000 members is ~14 MB of
#: HTML, and all 120,000 Silver members in one page would be ~300 MB -- a
#: request likely to time out and heavy on Zenith. If Zenith caps lower
#: (say 500), whatever it lists is used as it is.
LIMIT_TRY = 5000


def probe_page_limit(searcher, store: FFPStore, *, stop: threading.Event,
                    level: str = "Gold", log=lambda _m: None):
    """Search `level` with each candidate setting at 5000; keep the first
    that lists more rows than a plain search. Returns what it found (and
    stores it, so every later search uses it), or None."""
    base = searcher.search(level=level)
    store.save_members(base.members)
    shown = len(base.members)
    log(f"{level}: {base.total:,} in Zenith, {shown} listed by a plain "
        f"search. Trying larger pages…")
    if base.total <= shown:
        log(f"{level} is too small to tell -- pick a bigger level.")
        return None
    for where in ("form", "query"):
        for name in LIMIT_NAMES:
            if stop.is_set():
                return None
            page = searcher.search(level=level, **{
                f"extra_{where}": {name: str(LIMIT_TRY)}})
            if len(page.members) <= shown:
                continue
            store.save_members(page.members)
            found = {"where": where, "name": name, "value": LIMIT_TRY,
                     "listed": len(page.members)}
            store.put("page_param", json.dumps(found))
            log(f"Zenith listed {len(page.members):,} of {page.total:,} "
                f"with {name}={LIMIT_TRY} (sent in the {where}). Every "
                f"search will use it.")
            return found
    log("Zenith ignored every setting tried: 50 rows is a hard limit.")
    return None


def _close_small_levels(searcher, store: FFPStore, levels, log) -> None:
    """With a larger page, a whole level may fit in one search: take it
    and drop that level's queued name searches."""
    extra = _page_extra(store)
    if not extra:
        return
    for level in (levels or LEVEL_ORDER):
        if store.next_pending([level]) is None:
            continue
        page = searcher.search(level=level, **extra)
        if page.complete:
            store.save_members(page.members)
            store.close_level(level)
            log(f"{level}: all {page.total:,} in one search.")


def _too_deep(store: FFPStore, prefix: str) -> bool:
    if store.get("mode") == NUMBER:
        return len(prefix) >= MAX_DIGITS
    return len(prefix.split("|", 1)[1]) >= MAX_NAME


def _one_search(searcher, store: FFPStore, prefix: str, log):
    """Run one search and file what it found; split it if it was too big."""
    page = _search(searcher, store, prefix)
    if page.total == 0:
        store.mark(prefix, DONE, 0)
    elif page.complete:
        store.save_members(page.members)
        store.mark(prefix, DONE, page.total)
    elif _too_deep(store, prefix):
        store.save_members(page.members)
        store.mark(prefix, CAPPED, page.total)
        log(f"'{prefix}' still matches {page.total:,}; the first 50 "
            f"are kept.")
    else:
        store.save_members(page.members)
        store.seed(_children(store, prefix))
        store.mark(prefix, SPLIT, page.total)
    return page


def pending_levels(store: FFPStore) -> list:
    """Levels that still have searches to run, in collection order."""
    left = []
    for level in LEVEL_ORDER:
        if store.next_pending([level]) is not None:
            left.append(level)
    return left


def collect(searcher, store: FFPStore, *, stop: threading.Event,
            log=lambda _m: None, on_progress=lambda _p: None,
            delay_s: float = 0.5, workers: int = 1,
            levels=None) -> Progress:
    """Collect members into `store`, carrying on from any earlier run.

    `levels` limits the run to those levels (None: all), taken in
    LEVEL_ORDER -- the small levels first. `workers` searches run at once,
    each claiming its own, so waiting on Zenith overlaps. The first error
    stops every worker; the search it was on goes back in the queue.
    """
    if store.get("mode") not in (NUMBER, NAME):
        _choose_method(searcher, store, log)
    store.release_running()
    if store.get("mode") == NAME and store.get("space_fix") != "1":
        dropped = repair_space_chains(store)
        if dropped:
            log(f"Dropped {dropped:,} queued searches that only added "
                f"spaces (Zenith ignores a trailing space).")
    _close_small_levels(searcher, store, levels, log)
    errors: list = []

    def work() -> None:
        while not stop.is_set():
            prefix = store.claim_next(levels)
            if prefix is None:
                if store.running() == 0:
                    return
                stop.wait(0.3)       # another worker may still split one
                continue
            try:
                page = _one_search(searcher, store, prefix, log)
            except Exception as exc:          # noqa: BLE001 - handed back
                store.mark(prefix, PENDING, -1)
                errors.append(exc)
                stop.set()
                return
            on_progress(progress_of(store, searcher.searches,
                                    f"{prefix}: {page.total:,}"))
            if delay_s > 0:
                stop.wait(delay_s)

    pool = [threading.Thread(target=work, daemon=True)
            for _ in range(max(1, int(workers)))]
    for t in pool:
        t.start()
    for t in pool:
        t.join()
    store.release_running()
    store.put("finished", "1" if store.next_pending() is None else "0")
    if errors:
        raise errors[0]
    return progress_of(store, searcher.searches)


# ---- the workbook ---------------------------------------------------------

HEADERS = ("FFP number", "Customer ID", "Last name", "First name",
           "Date of birth", "Email", "Phone", "ID number", "Level",
           "Total miles")


def export_workbook(store: FFPStore, out_path) -> int:
    """Write every member found, and a summary against Zenith's own counts.

    Written row by row (write-only mode): the membership runs past 100,000.
    """
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    summary = wb.create_sheet("Summary")
    found = store.counts_by_level()
    totals = store.level_totals()
    summary.append(["FFP MEMBERS — CONFIDENTIAL: holds customers' personal "
                    "details. Do not forward outside the company."])
    summary.append([])
    summary.append(["Level", "In Zenith", "Collected", "Missing"])
    for level in zf.LEVELS:
        z, f = totals.get(level, -1), found.get(level, 0)
        summary.append([level, z if z >= 0 else None, f,
                        (z - f) if z >= 0 else None])
    others = {k: v for k, v in found.items() if k not in zf.LEVELS}
    for level, n in sorted(others.items()):
        summary.append([level or "(no level)", None, n, None])
    summary.append([])
    finished = store.get("finished") == "1"
    summary.append(["Collection", "complete" if finished else
                    "NOT FINISHED — press Collect to carry on"])
    gaps = store.capped()
    if gaps:
        summary.append(["Could not narrow", ", ".join(gaps)])

    sheet = wb.create_sheet("Members")
    sheet.append(list(HEADERS))
    n = 0
    for m in store.members():
        sheet.append([m[c] for c in _MEMBER_COLS])
        n += 1
    wb.save(str(out_path))
    return n


def default_filename() -> str:
    return f"FFP_Members_{time.strftime('%Y-%m-%d')}.xlsx"
