"""Zenith exchange rates (ROE): from the daily rate sheet to Zenith, checked.

The rate sheet (from the Travelport automation) gives "1 KWD = 401.00799"
BDT. Zenith keeps the reverse -- 1 BDT = x KWD -- with at most 6 decimals,
and the team CUTS to 6 decimals (1/401.00799 = 0.0024937 -> 0.002493):
proven against every currency of the 7 Oct sheet.

The ROE screen (SaisieTauxDevises.asp) lists the currencies over pages and
saves one page at a time by posting the page's whole form with the new
rates. This module only ever posts back the page it has just read, with
nothing changed but the approved rates (see RoePlan.permits), and checks
the result five ways (see verify).
"""
from __future__ import annotations

import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_DOWN, Decimal, InvalidOperation
from html import unescape

#: A change bigger than this (either way) is flagged for a second look.
BIG_CHANGE_PCT = Decimal("2")
#: A change bigger than this is not saved at all -- a typo on the sheet
#: (40.1 for 401) rather than a market move.
BLOCK_CHANGE_PCT = Decimal("10")
#: Quick Sale needs a route and a day to show a conversion; any will do.
QS_PROBE = "DACDXB"
SIX = Decimal("0.000001")


class RoeError(Exception):
    """The ROE screen or the sheet is not as expected; nothing was saved."""


class ConcurrentChange(RoeError):
    """A rate moved between the preview and the save."""


# ---- the rate sheet ----------------------------------------------------------

@dataclass
class SheetRates:
    day: date | None
    rates: dict = field(default_factory=dict)     # "KWD" -> Decimal BDT
    notes: dict = field(default_factory=dict)     # "USD" -> "06-OCT-26"
    problems: list = field(default_factory=list)


def _number(value) -> tuple:
    """401.00799 or '123.54\xa0(06-OCT-26)' -> (Decimal, note)."""
    if isinstance(value, (int, float)):
        return Decimal(str(value)), ""
    text = str(value or "").replace("\xa0", " ").strip()
    m = re.match(r"^([0-9][0-9,]*\.?[0-9]*)\s*(?:\((.*)\))?$", text)
    if not m:
        raise ValueError(text)
    return Decimal(m.group(1).replace(",", "")), (m.group(2) or "").strip()


def read_rate_sheet(path) -> SheetRates:
    """The first sheet: 'Current Date' above, then '1 XXX | rate' rows; the
    current day's columns (A-B) only."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True, read_only=True)
    try:
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    out = SheetRates(day=None)
    for i, r in enumerate(rows):
        first = r[0] if r else None
        if isinstance(first, str) and first.strip().lower() == "current date":
            nxt = rows[i + 1][0] if i + 1 < len(rows) and rows[i + 1] else None
            if isinstance(nxt, datetime):
                out.day = nxt.date()
            elif isinstance(nxt, date):
                out.day = nxt
        m = re.match(r"^\s*1\s+([A-Za-z]{3})\s*$", str(first or ""))
        if not m:
            continue
        code = m.group(1).upper()
        try:
            rate, note = _number(r[1] if len(r) > 1 else None)
            if rate <= 0:
                raise ValueError(rate)
        except (ValueError, InvalidOperation):
            out.problems.append(f"{code}: '{r[1] if len(r) > 1 else ''}' is "
                                f"not a rate.")
            continue
        if code in out.rates:
            out.problems.append(f"{code} is on the sheet twice; the first "
                                f"is used.")
            continue
        out.rates[code] = rate
        if note:
            out.notes[code] = note
    if not out.rates:
        raise RoeError("No '1 XXX' rate rows found on the sheet's first tab.")
    return out


def zenith_rate(bdt_per_unit: Decimal) -> Decimal:
    """1 XXX = 401.00799 BDT -> 1 BDT = 0.002493 XXX (cut, never rounded)."""
    return (Decimal(1) / bdt_per_unit).quantize(SIX, rounding=ROUND_DOWN)


def rate_text(rate: Decimal) -> str:
    """As Zenith writes it: 0.002490 -> '0.00249', 1 -> '1'."""
    text = f"{rate.quantize(SIX):f}".rstrip("0").rstrip(".")
    return text or "0"


# ---- the ROE screen ----------------------------------------------------------

def browser_form(html: str, name: str) -> list:
    """The (field, value) pairs a browser posts for form `name`: unticked
    boxes left out, selects at their selected option, and the values the
    page's own script sets after loading (theForm.X.value = '...')."""
    m = re.search(r'<form\b[^>]*name="' + re.escape(name)
                  + r'"[^>]*>(.*?)</form>', html or "", re.S | re.I)
    if not m:
        raise RoeError(f"The '{name}' form is not on the page.")
    body = re.sub(r"<!--.*?-->", "", m.group(1), flags=re.S)
    pairs: list = []
    for tag in re.findall(r"<input\b[^>]*>|<select\b.*?</select>|"
                          r"<textarea\b.*?</textarea>", body, re.S | re.I):
        head = re.match(r"<\w+\b[^>]*>", tag).group(0)
        bare = re.sub(r'"[^"]*"|\'[^\']*\'', '""', head)   # attrs w/o values
        fname = re.search(r'\bname="([^"]*)"', head)
        if not fname or re.search(r"\bdisabled\b", bare, re.I):
            continue
        if tag.lower().startswith("<select"):
            opts = re.findall(r"<option\b([^>]*)>", tag, re.I)
            chosen = next((o for o in opts if re.search(
                r"\bselected\b", re.sub(r'"[^"]*"', '""', o), re.I)),
                opts[0] if opts else "")
            v = re.search(r'\bvalue="([^"]*)"', chosen)
            pairs.append([fname.group(1), unescape(v.group(1)) if v else ""])
            continue
        if tag.lower().startswith("<textarea"):
            inner = re.search(r">(.*)</textarea>", tag, re.S | re.I)
            pairs.append([fname.group(1), unescape(inner.group(1))])
            continue
        typed = re.search(r'\btype="([^"]*)"', head, re.I)
        kind = typed.group(1).lower() if typed else "text"
        if kind in ("button", "submit", "image", "reset", "file"):
            continue
        v = re.search(r'\bvalue="([^"]*)"', head)
        value = unescape(v.group(1)) if v else ""
        if kind in ("checkbox", "radio"):
            if not re.search(r"\bchecked\b", bare, re.I):
                continue
            value = value or "on"
        pairs.append([fname.group(1), value])
    # a select's value can be set by the page's script after loading (the
    # language); only those are taken -- an assignment in a click handler
    # (Save sets hAction) must never leak into the form
    selects = set(re.findall(r'<select\b[^>]*\bname="([^"]*)"', body, re.I))
    for fld, value in re.findall(
            r"theForm\.(\w+)\.value\s*=\s*'([^']*)'", html):
        for p in pairs:
            if p[0] == fld and fld in selects:
                p[1] = value
    return [tuple(p) for p in pairs]


@dataclass(frozen=True)
class RoeRow:
    index: int                 # n in DeviseTaux{n} on its page
    code: str
    name: str
    rate: Decimal
    text: str                  # as on the page
    modified: str              # "08/10/2026 03:53:20"
    by: str
    history_id: int
    linked_taxes: tuple        # ("E5", "T6")


@dataclass
class RoePage:
    nav: int
    pages: list
    form: list                 # browser_form pairs
    rows: dict                 # code -> RoeRow


def _row_html(html: str, at: int) -> str:
    start = html.rfind("<tr", 0, at)
    end = html.find("</tr>", at)
    return html[start:end if end > 0 else len(html)]


def parse_roe_page(html: str) -> RoePage:
    form = browser_form(html, "frmDevise")
    values = dict(form)
    rows = {}
    for m in re.finditer(r'<input\b[^>]*name="DeviseTaux(\d+)"[^>]*>', html):
        n, tag = int(m.group(1)), m.group(0)
        code = values.get(f"CodeDevise{n}", "")
        if not code:
            raise RoeError(f"Rate row {n} has no currency code.")
        try:
            rate = Decimal(values.get(f"DeviseTaux{n}", "").replace(",", "."))
        except InvalidOperation as exc:
            raise RoeError(f"{code}: the rate on the page is not a "
                           f"number.") from exc
        row = _row_html(html, m.start())
        text = re.sub(r"\s+", " ", unescape(re.sub(
            r"<script.*?</script>|<[^>]+>", " ", row, flags=re.S))).strip()
        mod = re.search(r"Last modified (\S+ \S+)(?: GMT)? By (.*?) History",
                        text)
        name = re.search(r"Currency (.*?) Rate:", text)
        hist = re.search(r"DisplayHistory\((\d+)\)", row)
        linked = re.search(r'data-linked-taxes="([^"]*)"', tag)
        rows[code] = RoeRow(
            n, code, name.group(1).strip() if name else "", rate,
            values.get(f"DeviseTaux{n}", ""),
            mod.group(1) if mod else "", mod.group(2).strip() if mod else "",
            int(hist.group(1)) if hist else 0,
            tuple(t.strip() for t in linked.group(1).split(",") if t.strip())
            if linked else ())
    if not rows:
        raise RoeError("No currency rates on the ROE page.")
    pages = sorted({int(p) for p in re.findall(r"aff_page_result\((\d+)\)",
                                               html)} | {int(values.get(
                                                   "NAV", "1") or 1)})
    return RoePage(int(values.get("NAV", "1") or 1), pages, form, rows)


@dataclass(frozen=True)
class HistoryEntry:
    when: str
    by: str
    old: Decimal | None
    new: Decimal | None


def parse_history(html: str) -> list:
    """Newest first: 'Modified the … GMT | By … | Rate: old → Rate: new'."""
    out = []
    for row in re.findall(r"<tr\b[^>]*>(.*?)</tr>", html or "", re.S):
        cells = [re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", c)))
                 .strip() for c in re.findall(r"<td\b[^>]*>(.*?)</td>", row,
                                              re.S)]
        when = next((c for c in cells if c.startswith("Modified the")), "")
        if not when:
            continue
        by = next((c for c in cells if c.startswith("By ")), "")
        rates = [c[5:].strip() for c in cells if c.startswith("Rate:")]

        def dec(s):
            try:
                return Decimal(s.replace(",", "."))
            except InvalidOperation:
                return None
        out.append(HistoryEntry(when[len("Modified the"):].replace(
            "GMT", "").strip(), by[3:].strip(),
            dec(rates[0]) if rates else None,
            dec(rates[1]) if len(rates) > 1 else None))
    return out


QS_RATE_RE = re.compile(r"Currency:\s*([A-Z]{3})\s*\(1 BDT = ([0-9.]+) [A-Z]{3}\)")


def quick_sale_rate(html: str):
    """'Currency:AED(1 BDT = 0.02969600 AED)' -> ('AED', Decimal)."""
    text = unescape(re.sub(r"<[^>]+>", " ", html or "")).replace("\xa0", " ")
    m = QS_RATE_RE.search(re.sub(r"\s+", " ", text))
    return (m.group(1), Decimal(m.group(2))) if m else (None, None)


# ---- the plan: what would change ---------------------------------------------

@dataclass
class RoeChange:
    code: str
    nav: int
    index: int
    bdt_per_unit: Decimal
    old: Decimal
    new: Decimal
    linked_taxes: tuple
    note: str = ""

    @property
    def pct(self) -> Decimal:
        return ((self.new - self.old) / self.old * 100) if self.old else \
            Decimal(0)

    @property
    def big(self) -> bool:
        return abs(self.pct) > BIG_CHANGE_PCT


@dataclass
class RoePlan:
    sheet_day: date | None
    pages: dict                    # nav -> RoePage as read for the preview
    changes: list                  # [RoeChange]
    same: list                     # codes already at the sheet's rate
    skipped: list                  # (code, reason)
    warnings: list
    read_at: float = field(default_factory=time.time)

    def changes_on(self, nav: int) -> list:
        return [c for c in self.changes if c.nav == nav]

    def linked_taxes(self) -> list:
        return sorted({t for c in self.changes for t in c.linked_taxes})

    def save_form(self, page: RoePage, update_linked: bool) -> list:
        """The page's own form with the approved rates, as the browser
        posts it when Save is pressed."""
        mine = {f"DeviseTaux{c.index}": rate_text(c.new)
                for c in self.changes_on(page.nav)}
        linked = update_linked and any(c.linked_taxes
                                       for c in self.changes_on(page.nav))
        keys = {k for k, _v in page.form}
        if not {"hAction", "UpdateLinkedTaxes"} <= keys or not set(mine) <= keys:
            raise RoeError("The ROE form is not as recorded (no hAction / "
                           "UpdateLinkedTaxes / rate field); nothing saved.")
        out = []
        for k, v in page.form:
            if k == "hAction":
                v = "Save"
            elif k == "UpdateLinkedTaxes":
                v = "1" if linked else "0"
            elif k in mine:
                v = mine[k]
            out.append((k, v))
        return out

    def permits(self, page: RoePage, update_linked: bool):
        """For the client's guard: exactly this save, once, nothing else."""
        expected = sorted(self.save_form(page, update_linked))
        used = []

        def check(form) -> bool:
            got = list(form.items()) if isinstance(form, dict) else list(form)
            if used or sorted(got) != expected:
                return False
            used.append(True)
            return True
        return check


def make_plan(sheet: SheetRates, pages: dict, *, today: date) -> RoePlan:
    by_code = {code: (nav, row) for nav, p in pages.items()
               for code, row in p.rows.items()}
    changes, same, skipped, warnings = [], [], [], []
    if sheet.day is None:
        warnings.append("The sheet has no 'Current Date'.")
    elif sheet.day != today:
        warnings.append(f"The sheet is for {sheet.day:%d %b %Y}, not today "
                        f"({today:%d %b %Y}).")
    warnings.extend(sheet.problems)
    for code, bdt in sorted(sheet.rates.items()):
        if code not in by_code:
            skipped.append((code, "not a currency on Zenith's ROE screen"))
            continue
        nav, row = by_code[code]
        new = zenith_rate(bdt)
        if new <= 0:
            skipped.append((code, "rate too small for 6 decimals"))
            continue
        if new == row.rate:
            same.append(code)
            continue
        pct = (new - row.rate) / row.rate * 100 if row.rate else Decimal(0)
        if abs(pct) > BLOCK_CHANGE_PCT:
            skipped.append((code, f"would move {pct:+.1f}% ({rate_text(row.rate)}"
                                  f" → {rate_text(new)}) — over "
                                  f"{BLOCK_CHANGE_PCT}%, so NOT saved; check "
                                  f"the sheet (a typo?) and set it by hand"))
            continue
        changes.append(RoeChange(code, nav, row.index, bdt, row.rate, new,
                                 row.linked_taxes, sheet.notes.get(code, "")))
    for code in sorted(set(by_code) - set(sheet.rates)):
        skipped.append((code, "not on the sheet — left as it is"))
    for c in changes:
        if c.big:
            warnings.append(f"{c.code} moves {c.pct:+.2f}% "
                            f"({rate_text(c.old)} → {rate_text(c.new)}).")
        if c.note:
            warnings.append(f"{c.code}: the sheet notes '{c.note}'.")
    return RoePlan(sheet.day, pages, changes, same, skipped, warnings)


# ---- reading and saving ------------------------------------------------------

def read_all_pages(client) -> dict:
    first = parse_roe_page(client.roe_page())
    pages = {first.nav: first}
    for nav in first.pages:
        if nav in pages:
            continue
        form = [(k, (str(nav) if k == "NAV" else "" if k == "hAction" else v))
                for k, v in first.form]
        pages[nav] = parse_roe_page(client.roe_post(form))
    return pages


@dataclass
class Check:
    code: str
    old: str
    new: str
    saved: str = "—"           # 1. Zenith's reply
    reread: str = "—"          # 2. a fresh read of every page
    modified: str = "—"        # 3. 'Last modified' moved
    history: str = "—"         # 4. newest history entry is old → new
    quick_sale: str = "—"      # 5. Quick Sale converts at the new rate

    @property
    def ok(self) -> bool:
        return all(v.startswith("✓") for v in (
            self.saved, self.reread, self.modified, self.history))


@dataclass
class RoeResult:
    checks: list = field(default_factory=list)
    untouched_ok: bool = False
    untouched_notes: list = field(default_factory=list)
    saved_pages: list = field(default_factory=list)
    failed: str = ""
    update_linked: bool = True
    started: str = ""
    finished: str = ""


def apply(client, plan: RoePlan, *, update_linked: bool = True,
          today: date | None = None, progress=lambda _m: None) -> RoeResult:
    """Save the plan page by page, then verify. Stops before any page whose
    rates moved since the preview."""
    today = today or date.today()
    res = RoeResult(update_linked=update_linked,
                    started=time.strftime("%Y-%m-%d %H:%M:%S"))
    res.checks = [Check(c.code, rate_text(c.old), rate_text(c.new))
                  for c in plan.changes]
    check_of = {c.code: c for c in res.checks}
    navs = sorted({c.nav for c in plan.changes})
    for nav in navs:
        progress(f"Re-reading page {nav} before saving…")
        fresh = read_all_pages(client)[nav]
        moved = [code for code, row in plan.pages[nav].rows.items()
                 if code in fresh.rows and (fresh.rows[code].rate != row.rate
                                            or fresh.rows[code].index
                                            != row.index)]
        if moved or set(fresh.rows) != set(plan.pages[nav].rows):
            res.failed = (f"Page {nav} changed since the preview ("
                          + (", ".join(moved) or "different currencies")
                          + "). Nothing more was saved — preview again.")
            break
        form = plan.save_form(fresh, update_linked)
        progress(f"Saving page {nav}…")
        try:
            reply = parse_roe_page(client.roe_post(
                form, permit=plan.permits(fresh, update_linked)))
        except Exception as exc:                      # noqa: BLE001
            res.failed = (f"Saving page {nav} failed: {exc}. Check the "
                          f"rates in Zenith before trying again.")
            break
        res.saved_pages.append(nav)
        for c in plan.changes_on(nav):
            got = reply.rows.get(c.code)
            check_of[c.code].saved = (
                "✓ shows " + rate_text(got.rate) if got and got.rate == c.new
                else f"✗ shows {rate_text(got.rate) if got else 'nothing'}")
    _verify(client, plan, res, today, progress)
    res.finished = time.strftime("%Y-%m-%d %H:%M:%S")
    return res


def _verify(client, plan, res, today, progress) -> None:
    saved = {c.code for c in plan.changes if c.nav in res.saved_pages}
    check_of = {c.code: c for c in res.checks}
    progress("Checking: reading every page again…")
    after = read_all_pages(client)
    now = {code: row for p in after.values() for code, row in p.rows.items()}
    before = {code: row for p in plan.pages.values()
              for code, row in p.rows.items()}
    wanted = {c.code: c.new for c in plan.changes if c.code in saved}
    for code, row in before.items():
        if code not in now:
            res.untouched_notes.append(f"{code} is no longer on the screen.")
            continue
        if code in wanted:
            chk = check_of[code]
            chk.reread = ("✓ " + rate_text(now[code].rate)
                          if now[code].rate == wanted[code]
                          else f"✗ {rate_text(now[code].rate)}")
            chk.modified = ("✓ " + now[code].modified
                            if now[code].modified != row.modified
                            else "✗ unchanged")
        elif now[code].rate != row.rate:
            res.untouched_notes.append(
                f"{code} changed too ({rate_text(row.rate)} → "
                f"{rate_text(now[code].rate)}) though it was not asked for.")
    res.untouched_ok = not res.untouched_notes
    for code in sorted(wanted):
        chk, row = check_of[code], now.get(code)
        if row and row.history_id:
            progress(f"Checking: {code} history…")
            top = parse_history(client.roe_history(row.history_id))
            first = top[0] if top else None
            want_old = before[code].rate
            chk.history = ("✓ " + f"{rate_text(first.old)} → "
                           f"{rate_text(first.new)} by {first.by}"
                           if first and first.new == wanted[code]
                           and first.old == want_old
                           else "✗ newest entry: " + (
                               f"{first.old} → {first.new}" if first
                               else "none"))
        progress(f"Checking: Quick Sale in {code}…")
        day = (today + timedelta(days=7)).strftime("%d%b%y").upper()
        try:
            cur, rate = quick_sale_rate(client.quick_sale(
                f"FDI{QS_PROBE}/D{day}/R,NP,FC-{code}"))
        except Exception as exc:                      # noqa: BLE001
            chk.quick_sale = f"? could not ask ({exc})"
            continue
        chk.quick_sale = ("✓ 1 BDT = " + rate_text(rate) if rate is not None
                          and rate.quantize(SIX) == wanted[code]
                          else "? shows " + (rate_text(rate) if rate
                                             is not None else "no rate")
                          + " — may take a moment to refresh")
    for chk in res.checks:
        if chk.code not in saved:
            chk.saved = "✗ not saved"
